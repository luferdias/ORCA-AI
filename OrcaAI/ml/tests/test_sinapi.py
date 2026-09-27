"""Contratos contra perdas silenciosas nos ZIPs oficiais do SINAPI."""

import hashlib
import importlib
import io
import json
import xml.etree.ElementTree as ET
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook


def relatorio_zip(path, *, competencia="2026-08", es_col=19,
                  custos=(188.15, 22.0), editar=None, situacao=True):
    """Fixture real de XLSX, inclusive cache zero nas fórmulas HYPERLINK."""
    wb = Workbook()
    csd = wb.active
    csd.title = "CSD"
    ana = wb.create_sheet("Analítico")
    ano, mes = competencia.split("-")
    for sheet in (csd, ana):
        sheet["A3"] = "Mês de Referência:"
        sheet["B3"] = f"{mes}/{ano}"
    csd["A2"] = "RELATÓRIO DE CUSTOS DE COMPOSIÇÕES - ENCARGOS SOCIAIS SEM DESONERAÇÃO"
    csd.cell(9, es_col, "ES")
    csd.cell(10, es_col, "Custo (R$)")
    csd.cell(10, es_col + 1, "%AS")
    for col, value in enumerate(["Grupo", "Código da\nComposição", "Descrição", "Unidade"], 1):
        csd.cell(10, col, value)
    headers = ["Grupo", "Código da\nComposição", "Tipo Item", "Código do\nItem", "Descrição", "Unidade", "Coeficiente"]
    if situacao:
        headers.append("Situação")
    for col, value in enumerate(headers, 1):
        ana.cell(10, col, value)
    comps = [(104658, "Piso tátil", "M2"), (88316, "Servente com encargos", "H")]
    for index in range(2, len(custos)):
        comps.append((110000 + index, f"Serviço {index}", "M2"))
    row = 11
    for index, ((code, desc, unit), cost) in enumerate(zip(comps, custos)):
        cr = 11 + index
        values = ["Acessibilidade", f'=HYPERLINK("#"&CELL("address",OFFSET(Analítico!$B$1,MATCH({code},Analítico!$B:$B,0)-1,3)),{code})', desc, unit]
        for col, value in enumerate(values, 1):
            csd.cell(cr, col, value)
        csd.cell(cr, es_col, cost)
        csd.cell(cr, es_col + 1, 0.125)
        main = ["Acessibilidade", code, None, None, desc, unit, None]
        if situacao:
            main.append("COM CUSTO")
        for col, value in enumerate(main, 1):
            ana.cell(row, col, value)
        child = ["Acessibilidade", code, "COMPOSICAO" if index == 0 else "INSUMO",
                 88316 if index == 0 else 6111, "Servente com encargos" if index == 0 else "Mão de obra", "H", 1.279 if index == 0 else 1.0]
        if situacao:
            child.append("COM CUSTO" if index == 0 else "COM PREÇO")
        for col, value in enumerate(child, 1):
            ana.cell(row + 1, col, value)
        row += 2
    if editar:
        editar(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    rewritten = io.BytesIO()
    namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = ET.fromstring(data)
                for cell in xml.findall(".//s:c", namespace):
                    if cell.find("s:f", namespace) is not None:
                        cell.find("s:v", namespace).text = "0"
                data = ET.tostring(xml, encoding="utf-8")
            target.writestr(info, data)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(f"SINAPI_Referência_{ano}_{mes}.xlsx", rewritten.getvalue())
    return path


def ler(path, competencia="2026-08", publicado_em="2026-09-11"):
    return importlib.import_module("orca_ml.sinapi").ler_zip_sinapi(path, competencia, publicado_em)


def manifesto_zip(path, competencia="2026-08", publicado="2026-09-11",
                  nome="manifesto_coleta_regressao.json", alteracoes=None):
    data = {
        "competencia": competencia, "arquivo": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "tamanho_bytes": path.stat().st_size,
        "publicacao_exibida_no_portal": publicado,
    }
    data.update(alteracoes or {})
    (path.parent / nome).write_text(json.dumps(data), encoding="utf-8")
    return data


def test_recupera_codigo_formula_com_cache_zero_e_valida_analitico(tmp_path):
    path = relatorio_zip(tmp_path / "sinapi.zip")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    with zipfile.ZipFile(path) as archive:
        wb = load_workbook(io.BytesIO(archive.read(archive.namelist()[0])), data_only=True)
        assert wb["CSD"]["B11"].value == 0
        wb.close()
    df = ler(path)
    assert df["codigo"].tolist() == ["104658", "88316"]
    assert df["custo"].tolist() == [188.15, 22.0]
    row = df.iloc[0]
    assert row["fonte"] == "sinapi_es"
    assert row["regime"] == "sem_desoneracao"
    assert row["unidade"] == "M2"
    assert row["grupo"] == "Acessibilidade"
    assert row["percentual_as"] == 0.125
    assert row["publicado_em"] == "2026-09-11"
    assert row["arquivo_sha256"] == before
    assert "CSD!S11" in row["localizador"]
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_descobre_coluna_es_pelos_cabecalhos_sem_posicao_fixa(tmp_path):
    def editar(wb):
        wb["CSD"]["S11"] = 9999
    df = ler(relatorio_zip(tmp_path / "outra_coluna.zip", es_col=7, editar=editar))
    assert df.iloc[0]["custo"] == 188.15
    assert "CSD!G11" in df.iloc[0]["localizador"]


@pytest.mark.parametrize("mutacao", [
    lambda w: setattr(w["CSD"]["S9"], "value", "SP"),
    lambda w: setattr(w["CSD"]["S10"], "value", "Preço errado"),
    lambda w: setattr(w["CSD"]["B3"], "value", "07/2026"),
    lambda w: setattr(w["CSD"]["A2"], "value", "RELATÓRIO COM DESONERAÇÃO"),
    lambda w: setattr(w["Analítico"]["G10"], "value", "Quantidade desconhecida"),
])
def test_falha_explicitamente_em_esquema_regime_ou_competencia_incompativel(tmp_path, mutacao):
    with pytest.raises(ValueError):
        ler(relatorio_zip(tmp_path / "schema.zip", editar=mutacao))


def test_exclui_hifen_zero_e_ausencia_com_relatorio_contavel(tmp_path):
    df = ler(relatorio_zip(tmp_path / "faltantes.zip", custos=(188.15, "-", 0, None)))
    assert df["codigo"].tolist() == ["104658"]
    report = df.attrs["relatorio"]
    assert report["composicoes_lidas"] == 4
    assert report["composicoes_validas"] == 1
    assert report["excluidos_por_motivo"] == {"custo_ausente": 2, "custo_zero": 1}
    assert len(report["exclusoes"]) == 3
    json.dumps(report, allow_nan=False)


def test_assinatura_estavel_quando_so_preco_status_ou_grupo_muda(tmp_path):
    original = ler(relatorio_zip(tmp_path / "a.zip"))
    def editar(wb):
        for row in wb["Analítico"].iter_rows(min_row=11):
            row[0].value = "Outro grupo"
            row[7].value = "SEM CUSTO"
        wb["CSD"]["A11"] = "Outro grupo"
        wb["CSD"]["A12"] = "Outro grupo"
    atualizado = ler(relatorio_zip(tmp_path / "b.zip", custos=(200, 28), editar=editar))
    assert original["assinatura_tecnica"].tolist() == atualizado["assinatura_tecnica"].tolist()


@pytest.mark.parametrize("celula,novo", [("G12", 1.5), ("G14", 2.0), ("D14", 7777), ("E14", "Outra mão de obra")])
def test_assinatura_propaga_coeficiente_ou_item_alterado_no_auxiliar(tmp_path, celula, novo):
    a = ler(relatorio_zip(tmp_path / "a.zip"))
    def editar(wb):
        wb["Analítico"][celula] = novo
    b = ler(relatorio_zip(tmp_path / "b.zip", editar=editar))
    assert a.iloc[0]["assinatura_tecnica"] != b.iloc[0]["assinatura_tecnica"]


def test_aceita_analitico_antigo_sem_coluna_situacao(tmp_path):
    a = ler(relatorio_zip(tmp_path / "a.zip"))
    b = ler(relatorio_zip(tmp_path / "b.zip", situacao=False))
    assert a["assinatura_tecnica"].tolist() == b["assinatura_tecnica"].tolist()


@pytest.mark.parametrize("aba", ["CSD", "Analítico"])
def test_rejeita_codigos_duplicados_antes_de_filtrar_custos(tmp_path, aba):
    def editar(wb):
        sheet = wb[aba]
        sheet.append([cell.value for cell in sheet[11]])
    with pytest.raises(ValueError, match="duplicad"):
        ler(relatorio_zip(tmp_path / "duplicados.zip", editar=editar))


@pytest.mark.parametrize("mutacao", [
    lambda w: setattr(w["CSD"]["B11"], "value", '=HYPERLINK("#",999999)'),
    lambda w: setattr(w["CSD"]["B11"], "value", '=HYPERLINK("#"&MATCH(88316,Analítico!$B:$B,0),104658)'),
    lambda w: setattr(w["CSD"]["C11"], "value", "Descrição sem correspondência"),
])
def test_rejeita_formula_codigo_ou_correspondencia_incorreta(tmp_path, mutacao):
    with pytest.raises(ValueError):
        ler(relatorio_zip(tmp_path / "id_errado.zip", editar=mutacao))


def test_rejeita_ciclo_de_composicoes(tmp_path):
    def editar(wb):
        wb["Analítico"]["C14"] = "COMPOSICAO"
        wb["Analítico"]["D14"] = 104658
    with pytest.raises(ValueError, match="[Cc]iclo"):
        ler(relatorio_zip(tmp_path / "ciclo.zip", editar=editar))


def test_rejeita_auxiliar_ausente_na_hierarquia(tmp_path):
    def editar(wb):
        wb["Analítico"]["D12"] = 999999
    with pytest.raises(ValueError, match="ausente"):
        ler(relatorio_zip(tmp_path / "auxiliar.zip", editar=editar))


def test_historico_grava_csv_e_relatorio_sem_alterar_originais(tmp_path):
    source = tmp_path / "fontes"
    for competencia, publicado in [("2026-07", "2026-08-11"), ("2026-08", "2026-09-11")]:
        originais = source / competencia / "originais"
        originais.mkdir(parents=True)
        path = relatorio_zip(originais / f"SINAPI-{competencia}-formato-xlsx.zip", competencia=competencia)
        manifesto_zip(path, competencia, publicado)
    output = tmp_path / "resultado" / "sinapi.csv"
    module = importlib.import_module("orca_ml.sinapi")
    df = module.preparar_historico_sinapi(source, output)
    assert len(df) == 4
    saved = pd.read_csv(output, dtype={"codigo": str})
    assert set(saved["codigo"]) == {"104658", "88316"}
    assert set(saved["publicado_em"]) == {"2026-08-11", "2026-09-11"}
    report = json.loads(output.with_suffix(".relatorio.json").read_text())
    assert report["composicoes_validas"] == 4
    assert len(report["edicoes"]) == 2


def test_historico_rejeita_duas_revisoes_da_mesma_competencia(tmp_path):
    originais = tmp_path / "fontes" / "2026-08" / "originais"
    originais.mkdir(parents=True)
    relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx.zip")
    relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx_Retificacao01.zip")
    module = importlib.import_module("orca_ml.sinapi")
    with pytest.raises(ValueError, match="revis|ediç|duplicad"):
        module.preparar_historico_sinapi(tmp_path / "fontes", tmp_path / "saida.csv")


def test_historico_prioriza_publicacao_observada_no_manifesto_de_coleta(tmp_path):
    originais = tmp_path / "fontes" / "2026-08" / "originais"
    originais.mkdir(parents=True)
    path = relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx.zip")
    manifesto_zip(path, publicado="2026-09-15")
    module = importlib.import_module("orca_ml.sinapi")
    df = module.preparar_historico_sinapi(tmp_path / "fontes", tmp_path / "saida.csv")
    assert df["publicado_em"].unique().tolist() == ["2026-09-15"]


def test_historico_rejeita_zip_trocado_mesmo_com_crc_valido(tmp_path):
    originais = tmp_path / "fontes" / "2026-08" / "originais"
    originais.mkdir(parents=True)
    path = relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx.zip")
    manifesto_zip(path)
    relatorio_zip(path, custos=(999, 888))
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    output = tmp_path / "saida.csv"
    output.write_text("resultado anterior")
    module = importlib.import_module("orca_ml.sinapi")
    with pytest.raises(ValueError, match="SHA|tamanho|bytes"):
        module.preparar_historico_sinapi(tmp_path / "fontes", output)
    assert output.read_text() == "resultado anterior"
    assert not output.with_suffix(".relatorio.json").exists()


@pytest.mark.parametrize("alteracoes", [
    {"arquivo": "SINAPI-2026-08-formato-xlsx_Retificacao01.zip"},
    {"sha256": "0" * 64},
    {"sha256": None},
    {"tamanho_bytes": 1},
    {"tamanho_bytes": None},
    {"competencia": "2026-07"},
    {"competencia": None},
])
def test_historico_rejeita_metadados_de_identidade_divergentes(tmp_path, alteracoes):
    originais = tmp_path / "fontes" / "2026-08" / "originais"
    originais.mkdir(parents=True)
    path = relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx.zip")
    manifesto_zip(path, alteracoes=alteracoes)
    module = importlib.import_module("orca_ml.sinapi")
    with pytest.raises(ValueError):
        module.preparar_historico_sinapi(tmp_path / "fontes", tmp_path / "saida.csv")


def test_catalogo_de_datas_nao_substitui_manifesto_de_download(tmp_path):
    originais = tmp_path / "fontes" / "2026-08" / "originais"
    originais.mkdir(parents=True)
    relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx.zip")
    module = importlib.import_module("orca_ml.sinapi")
    with pytest.raises(ValueError, match="[Mm]anifesto"):
        module.preparar_historico_sinapi(tmp_path / "fontes", tmp_path / "saida.csv")


@pytest.mark.parametrize("alteracoes", [
    {"arquivo": "outra_revisao.zip"}, {"sha256": "0" * 64},
    {"competencia": "2026-07"}, {"publicacao_exibida_no_portal": "2026-09-12"},
])
def test_confere_todos_os_manifestos_presentes_sem_ocultar_divergencia(tmp_path, alteracoes):
    originais = tmp_path / "fontes" / "2026-08" / "originais"
    originais.mkdir(parents=True)
    path = relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx.zip")
    manifesto_zip(path)
    manifesto_zip(path, nome="manifesto_download.json", alteracoes=alteracoes)
    module = importlib.import_module("orca_ml.sinapi")
    with pytest.raises(ValueError):
        module.preparar_historico_sinapi(tmp_path / "fontes", tmp_path / "saida.csv")


def test_aceita_bytes_e_dois_manifestos_concordantes(tmp_path):
    originais = tmp_path / "fontes" / "2026-08" / "originais"
    originais.mkdir(parents=True)
    path = relatorio_zip(originais / "SINAPI-2026-08-formato-xlsx.zip")
    manifesto_zip(path)
    data = manifesto_zip(path, nome="manifesto_download.json")
    data["bytes"] = data.pop("tamanho_bytes")
    (originais / "manifesto_download.json").write_text(json.dumps(data))
    module = importlib.import_module("orca_ml.sinapi")
    df = module.preparar_historico_sinapi(tmp_path / "fontes", tmp_path / "saida.csv")
    assert len(df) == 2
    assert df["publicado_em"].unique().tolist() == ["2026-09-11"]
