"""Contratos do importador dos relatórios oficiais DER-ES de edificações."""

import csv
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from orca_ml.der import COLUNAS, extrair_links_oficiais, parse_der_zip, preparar_csv


def zip_relatorio(path, *, titulo="SERVIÇOS NÃO DESONERADOS MAIO/2026", mes="Maio/2026", preco=18.12):
    wb = Workbook()
    ws = wb.active
    ws.title = "relatorio(166)"
    ws["A6"] = titulo
    ws["A7"] = "BDI: 0,00%"
    ws["A9"] = f"Data Base: {mes}"
    for col, value in enumerate(["Item", "Fonte/Código", "Especificação do Serviço", "Und.", "Quant.", "Preço Unitário", "Preço Total"], 1):
        ws.cell(12, col, value)
    ws.append(["'01", None, "SERVIÇOS PRELIMINARES", None, None, None, None])
    ws.append(["'0102", None, "DEMOLIÇÕES E RETIRADAS", None, None, None, None])
    ws.append(["'010201", "LABOR - 2024 - 010201 - 2", "Demolição manual de piso", "m2", 1, preco, preco])
    ws.append([None, None, "TOTAL", None, None, 9999, None])
    buffer = io.BytesIO()
    wb.save(buffer)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("pasta/XLS/tab_DER-EDIFICAÇÕES_2026_05_servicos.xlsx", buffer.getvalue())


class TestDer(unittest.TestCase):
    def test_parser_preserva_codigo_grupo_preco_e_assinatura_estavel(self):
        with tempfile.TemporaryDirectory() as temp:
            primeiro, segundo = Path(temp) / "a.zip", Path(temp) / "b.zip"
            zip_relatorio(primeiro)
            zip_relatorio(segundo, preco=21.30)
            linha = parse_der_zip(primeiro, "2026-05")[0]
            outra = parse_der_zip(segundo, "2026-05")[0]
            self.assertEqual(len(parse_der_zip(primeiro, "2026-05")), 1)
            self.assertEqual(linha["codigo"], "010201")
            self.assertEqual(linha["grupo"], "SERVIÇOS PRELIMINARES / DEMOLIÇÕES E RETIRADAS")
            self.assertEqual(linha["custo"], 18.12)
            self.assertEqual(linha["regime"], "sem_desoneracao")
            self.assertEqual(linha["assinatura_tipo"], "descricao_unidade")
            self.assertEqual(linha["assinatura_tecnica"], outra["assinatura_tecnica"])
            self.assertNotEqual(linha["arquivo_sha256"], outra["arquivo_sha256"])
            self.assertIn("servicos.xlsx!A15", linha["localizador"])

    def test_rejeita_regime_e_competencia_inconsistentes(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "a.zip"
            zip_relatorio(path, titulo="SERVIÇOS DESONERADOS MAIO/2026")
            with self.assertRaisesRegex(ValueError, "regime"):
                parse_der_zip(path, "2026-05")
            zip_relatorio(path)
            with self.assertRaisesRegex(ValueError, "competência"):
                parse_der_zip(path, "2026-04")

    def test_aceita_titulo_antigo_sem_palavra_servicos(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "a.zip"
            zip_relatorio(path, titulo="TABELA DER-ES NÃO DESONERADOS MAIO/2026")
            self.assertEqual(len(parse_der_zip(path, "2026-05")), 1)

    def test_links_do_portal_vem_de_href_observado(self):
        html = '''<a href="/media/Documentos/Tabela%20Referencial%20de%20Pre%C3%A7o/Referencial%20de%20Edifica%C3%A7%C3%B5es/OBRAS_2026/TABELA%20REFERENCIAL%20OBRAS%20-%20DER-ES%20-%202026-05%20-%20MAI.zip">MAI/2026</a>
        <a href="/media/Documentos/rodovias/2026-05.zip">rodovias</a>'''
        links = extrair_links_oficiais(html)
        self.assertEqual(list(links), ["2026-05"])
        self.assertTrue(links["2026-05"].startswith("https://der.es.gov.br/media/"))

    def test_csv_tem_esquema_exato_e_rejeita_revisao_duplicada(self):
        with tempfile.TemporaryDirectory() as temp:
            path, saida = Path(temp) / "a.zip", Path(temp) / "serie.csv"
            zip_relatorio(path)
            rows = parse_der_zip(path, "2026-05")
            preparar_csv(rows, saida)
            with saida.open(encoding="utf-8", newline="") as stream:
                reader = csv.DictReader(stream)
                self.assertEqual(reader.fieldnames, COLUNAS)
                self.assertEqual(next(reader)["publicado_em"], "")
            with self.assertRaisesRegex(ValueError, "duplicad"):
                preparar_csv(rows + rows, saida)


if __name__ == "__main__":
    unittest.main()
