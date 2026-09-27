"""Falhas de extração não podem virar perfis aparentemente completos."""

import hashlib
import importlib
import io
import json
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook


def pacote(tmp_path, editar=None):
    wb = Workbook()
    csd = wb.active
    csd.title = 'CSD'
    ana = wb.create_sheet('Analítico')
    ins = wb.create_sheet('ISD')
    for sheet in (csd, ana, ins):
        sheet['B3'] = '05/2026'
    csd['A2'] = 'RELATÓRIO SEM DESONERAÇÃO'
    ins['A2'] = 'RELATÓRIO SEM DESONERAÇÃO'
    for c, v in enumerate(['Grupo', 'Código da Composição', 'Descrição', 'Unidade'], 1):
        csd.cell(10, c, v)
    csd['S9'], csd['S10'], csd['T10'] = 'ES', 'Custo (R$)', '%AS'
    for c, v in enumerate(['Grupo', 'Código da Composição', 'Tipo Item', 'Código do Item', 'Descrição', 'Unidade', 'Coeficiente', 'Situação'], 1):
        ana.cell(10, c, v)
    for c, v in enumerate(['Classificação', 'Código do Insumo', 'Descrição do Insumo', 'Unidade', 'Origem de Preço'], 1):
        ins.cell(10, c, v)
    ins['M10'], ins['AE10'] = 'ES', 'SP'
    for row in [
        ['Pintura Interna', 100, None, None, 'Pintura exemplo', 'M2', None],
        ['Pintura Interna', 100, 'COMPOSICAO', 200, 'Auxiliar', 'H', 2],
        ['Pintura Interna', 100, 'INSUMO', 3, 'Tinta', 'L', 1.3],
        ['Pintura Interna', 100, 'INSUMO', 4, 'Equipamento', 'H', .5],
        ['Mão de obra', 200, None, None, 'Auxiliar', 'H', None],
        ['Mão de obra', 200, 'INSUMO', 1, 'Pintor', 'H', .333],
        ['Mão de obra', 200, 'INSUMO', 2, 'EPI', 'H', .333],
        ['Mão de obra', 200, 'INSUMO', 3, 'Tinta', 'L', .155],
    ]:
        ana.append(row)
    for code, desc, unit, cost in [(100, 'Pintura exemplo', 'M2', 8.49), (200, 'Auxiliar', 'H', 1.64)]:
        csd.append(['Pintura Interna' if code == 100 else 'Mão de obra', f'=HYPERLINK("#",{code})', desc, unit] + [None] * 14 + [cost, 0])
    for code, category, desc, unit, es, sp in [
        (1, 'MAO DE OBRA', 'Pintor', 'H', 3., 5.),
        (2, 'ENCARGOS COMPLEMENTARES', 'EPI', 'H', .1, .2),
        (3, 'MATERIAL', 'Tinta', 'L', 4., 4.),
        (4, 'EQUIPAMENTO (LOCAÇÃO)', 'Equipamento', 'H', .039, .04),
    ]:
        row = ins.max_row + 1
        for c, value in enumerate([category, code, desc, unit, 'C'], 1):
            ins.cell(row, c, value)
        ins.cell(row, 13, es)
        ins.cell(row, 31, sp)
    if editar:
        editar(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    path = tmp_path / 'SINAPI.zip'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('SINAPI_Referência_2026_05.xlsx', buffer.getvalue())
    return path


def snapshot(path):
    return pd.DataFrame([{
        'fonte': 'sinapi_es', 'competencia': '2026-05', 'codigo': '100',
        'descricao': 'Pintura exemplo', 'grupo': 'Pintura Interna', 'unidade': 'M2',
        'custo': 8.49, 'regime': 'sem_desoneracao', 'percentual_as': 0.,
        'arquivo_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
    }])


def ler(path, servicos=None):
    module = importlib.import_module('orca_ml.sinapi_composicoes')
    return module.ler_composicoes_sinapi(path, '2026-05', snapshot(path) if servicos is None else servicos)


def test_expande_subcomposicoes_agrega_folhas_e_concilia_por_nivel(tmp_path):
    path = pacote(tmp_path)
    before = path.read_bytes()
    services, resources, audit = ler(path)
    assert services.status_composicao.tolist() == ['ok']
    assert services.custo_recalculado_hierarquia.tolist() == [8.49]
    assert services.diferenca_reconciliacao.tolist() == [0.]
    assert len(resources) == 4  # Tinta direta e indireta não duplicam o recurso.
    by_code = resources.set_index('codigo_recurso')
    assert by_code.loc['3', 'coeficiente'] == pytest.approx(1.61)
    assert by_code.loc['3', 'custo_folha_sem_truncamento'] == pytest.approx(6.44)
    assert by_code.loc['1', 'horas_mao_obra'] == pytest.approx(.666)
    assert by_code.loc['2', 'horas_mao_obra'] == 0.
    assert by_code.loc['2', 'categoria'] == 'mao_obra'
    assert by_code.loc['4', 'horas_equipamento'] == .5
    assert 'ISD!M13' in by_code.loc['3', 'localizador']
    assert 'Analítico!' in by_code.loc['3', 'localizador']
    assert json.loads(json.dumps(audit, allow_nan=False))['servicos_validos'] == 1
    assert path.read_bytes() == before


def test_fallback_sp_preserva_ausencia_es_e_confere_indicador_oficial(tmp_path):
    def editar(w):
        w['ISD']['M13'] = None
        w['CSD']['T11'] = .7585
        w['CSD']['T12'] = .3780
    path = pacote(tmp_path, editar)
    selected = snapshot(path)
    selected['percentual_as'] = .7585
    services, resources, audit = ler(path, selected)
    assert services.status_composicao.tolist() == ['ok']
    material = resources.set_index('codigo_recurso').loc['3']
    assert pd.isna(material.preco_es)
    assert material.preco_efetivo == 4.
    assert material.uf_preco == 'SP'
    assert 'ISD!AE13' in material.localizador
    assert audit['servicos_com_fallback_sp'] == 1


def test_sp_sem_indicador_as_exclui_sem_disfarcar_perfil_incompleto(tmp_path):
    path = pacote(tmp_path, lambda w: setattr(w['ISD']['M13'], 'value', None))
    services, resources, audit = ler(path)
    assert services.status_composicao.tolist() == ['excluido']
    assert 'percentual_as' in services.iloc[0].motivo_exclusao
    assert resources.empty
    assert audit['servicos_excluidos'] == 1


@pytest.mark.parametrize('mutacao,motivo', [
    (lambda w: setattr(w['ISD']['M13'], 'value', '-'), 'sem_preco'),
    (lambda w: setattr(w['Analítico']['D12'], 'value', 999), 'auxiliar_ausente'),
    (lambda w: setattr(w['Analítico']['D12'], 'value', 100), 'ciclo'),
    (lambda w: setattr(w['ISD']['A13'], 'value', 'CATEGORIA NOVA'), 'categoria_desconhecida'),
    (lambda w: setattr(w['ISD']['D13'], 'value', 'KG'), 'unidade_divergente'),
    (lambda w: setattr(w['ISD']['M13'], 'value', 5.), 'custo_nao_conciliado'),
])
def test_problema_de_composicao_preserva_servico_com_exclusao(tmp_path, mutacao, motivo):
    def editar(w):
        mutacao(w)
        if motivo == 'sem_preco':
            w['ISD']['AE13'] = None
    path = pacote(tmp_path, editar)
    services, resources, audit = ler(path)
    assert services.codigo.tolist() == ['100']
    assert services.status_composicao.tolist() == ['excluido']
    assert motivo in services.iloc[0].motivo_exclusao
    assert resources.empty
    assert audit['exclusoes'][0]['codigo'] == '100'


def test_codigo_ausente_na_fonte_e_registrado_sem_descartar_linha(tmp_path):
    path = pacote(tmp_path)
    selected = snapshot(path)
    selected.loc[1] = dict(selected.iloc[0], codigo='999')
    services, resources, _ = ler(path, selected)
    assert services.codigo.tolist() == ['100', '999']
    assert services.status_composicao.tolist() == ['ok', 'excluido']
    assert 'composicao_ausente' in services.iloc[1].motivo_exclusao
    assert set(resources.codigo_servico) == {'100'}


@pytest.mark.parametrize('campo,valor', [('custo', 10.), ('descricao', 'Outra'), ('grupo', 'Outro'), ('unidade', 'M')])
def test_snapshot_divergente_da_planilha_exclui_composicao(tmp_path, campo, valor):
    path = pacote(tmp_path)
    selected = snapshot(path)
    selected[campo] = valor
    services, resources, _ = ler(path, selected)
    assert services.status_composicao.tolist() == ['excluido']
    assert 'snapshot_divergente' in services.iloc[0].motivo_exclusao
    assert resources.empty


def test_rejeita_dataset_de_outra_edicao_antes_da_extracao(tmp_path):
    path = pacote(tmp_path)
    selected = snapshot(path)
    selected['competencia'] = '2026-04'
    with pytest.raises(ValueError, match='competencia'):
        ler(path, selected)
    selected['competencia'] = '2026-05'
    selected['arquivo_sha256'] = '0' * 64
    with pytest.raises(ValueError, match='SHA'):
        ler(path, selected)


def test_rejeita_duplicata_de_servico_e_categoria_negativa_nao_vira_zero(tmp_path):
    path = pacote(tmp_path)
    selected = pd.concat([snapshot(path), snapshot(path)], ignore_index=True)
    with pytest.raises(ValueError, match='duplicad'):
        ler(path, selected)
    path = pacote(tmp_path, lambda w: setattr(w['ISD']['M13'], 'value', -1.))
    with pytest.raises(ValueError, match='negativo'):
        ler(path)


def test_preco_zero_publicado_nao_e_trocado_por_preco_sp(tmp_path):
    def editar(w):
        w['ISD']['M14'] = 0
        w['CSD']['S11'] = 8.48
    path = pacote(tmp_path, editar)
    selected = snapshot(path)
    selected['custo'] = 8.48
    services, resources, audit = ler(path, selected)
    assert services.status_composicao.tolist() == ['ok']
    equipment = resources.set_index('codigo_recurso').loc['4']
    assert equipment.preco_efetivo == 0
    assert equipment.uf_preco == 'ES'
    assert equipment.coeficiente == .5
    assert audit['servicos_com_fallback_sp'] == 0


def test_insumo_ausente_exclui_todo_perfil_e_nao_so_a_linha(tmp_path):
    path = pacote(tmp_path, lambda w: setattr(w['Analítico']['D13'], 'value', 999))
    services, resources, audit = ler(path)
    assert services.status_composicao.tolist() == ['excluido']
    assert 'insumo_ausente' in services.iloc[0].motivo_exclusao
    assert resources.empty


def test_percentual_as_positivo_sem_folha_sp_e_inconsistente(tmp_path):
    path = pacote(tmp_path, lambda w: setattr(w['CSD']['T11'], 'value', .5))
    selected = snapshot(path)
    selected['percentual_as'] = .5
    services, resources, _ = ler(path, selected)
    assert services.status_composicao.tolist() == ['excluido']
    assert 'percentual_as' in services.iloc[0].motivo_exclusao
    assert resources.empty


def test_duplicata_isd_impede_escolha_arbitraria_de_preco(tmp_path):
    def editar(w):
        w['ISD'].append([cell.value for cell in w['ISD'][13]])
    path = pacote(tmp_path, editar)
    with pytest.raises(ValueError, match='duplicado'):
        ler(path)
