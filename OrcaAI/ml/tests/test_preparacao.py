"""Reconstrução offline: arquivos originais, hashes e edição única."""

import csv
import hashlib
import importlib
import io
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

ML = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ML))


def fixture_der(base):
    directory = Path(base) / "2026-05" / "originais"
    directory.mkdir(parents=True)
    archive = directory / "tabela maio.zip"
    wb = Workbook()
    ws = wb.active
    ws["A6"] = "SERVIÇOS NÃO DESONERADOS MAIO/2026"
    ws["A7"] = "BDI: 0,00%"
    ws["A9"] = "Data Base: Maio/2026"
    for i, v in enumerate(["Item", "Fonte/Código", "Especificação do Serviço", "Und.", "Quant.", "Preço Unitário", "Preço Total"], 1):
        ws.cell(12, i, v)
    ws.append(["'01", None, "SERVIÇOS PRELIMINARES", None, None, None, None])
    ws.append(["'010201", "LABOR - 2024", "Demolição de piso", "m2", 1, 18.12, 18.12])
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("relatorios/XLS/tabela_servicos.xlsx", buffer.getvalue())
    manifest = directory / "manifesto.json"
    manifest.write_text(json.dumps({
        "arquivo": archive.name, "sha256": hashlib.sha256(archive.read_bytes()).hexdigest(),
        "bytes": archive.stat().st_size, "fonte": "DER-ES", "competencia": "2026-05",
        "regime": "sem_desoneracao", "publicado_em": None,
    }), encoding="utf-8")
    return archive, manifest


class PreparacaoTest(unittest.TestCase):
    def setUp(self):
        self.assertTrue((ML / "preparar_dados.py").exists(), "Falta reconstrução offline real")
        self.mod = importlib.import_module("preparar_dados")
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name) / "base"
        self.archive, self.manifest = fixture_der(self.base)

    def test_manifesto_confere_nome_exato_tamanho_e_hash(self):
        manifest, archive = self.mod.conferir_manifesto(self.manifest)
        self.assertEqual(archive, self.archive)
        self.assertEqual(manifest["sha256"], hashlib.sha256(archive.read_bytes()).hexdigest())

    def test_hash_divergente_interrompe_antes_de_sobrescrever_dataset(self):
        self.archive.write_bytes(self.archive.read_bytes() + b"alteracao")
        destination = Path(self.temp.name) / "existente.csv"
        destination.write_text("anterior\n")
        with self.assertRaisesRegex(ValueError, "SHA-256"):
            self.mod.preparar_der(self.base, destination)
        self.assertEqual(destination.read_text(), "anterior\n")

    def test_nao_resolve_nome_por_aproximacao_nem_aceita_escape(self):
        manifest = json.loads(self.manifest.read_text())
        for name in ["outro.zip", "../tabela maio.zip"]:
            manifest["arquivo"] = name
            self.manifest.write_text(json.dumps(manifest))
            with self.subTest(name=name), self.assertRaises((ValueError, FileNotFoundError)):
                self.mod.conferir_manifesto(self.manifest)

    def test_revisoes_ambiguas_falham(self):
        (self.archive.parent / "segunda revisao.zip").write_bytes(self.archive.read_bytes())
        with self.assertRaisesRegex(ValueError, "revis|ZIP"):
            self.mod.preparar_der(self.base, Path(self.temp.name) / "saida.csv")

    def test_reconstroi_der_com_data_ausente_e_curadoria_sem_metrica_ml(self):
        before = self.archive.read_bytes(), self.manifest.read_bytes()
        destination = Path(self.temp.name) / "der_es" / "custos_historicos.csv"
        report = self.mod.preparar_der(self.base, destination)
        with destination.open(newline="", encoding="utf-8") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["codigo"], "010201")
        self.assertEqual(rows[0]["publicado_em"], "")
        self.assertEqual(float(rows[0]["custo"]), 18.12)
        self.assertEqual(report["registros"], 1)
        self.assertEqual(report["competencias"][0]["codigos_unicos"], 1)
        self.assertEqual(report["competencias"][0]["custos_por_unidade"]["m2"]["minimo"], 18.12)
        self.assertEqual(report["assinatura_tipo"], "descricao_unidade")
        self.assertFalse(report["metricas_de_modelo_calculadas"])
        self.assertEqual(json.loads((destination.parent / "curadoria.json").read_text()), report)
        self.assertEqual((self.archive.read_bytes(), self.manifest.read_bytes()), before)

    def test_ajuda_cli_independe_do_diretorio_atual(self):
        result = subprocess.run([sys.executable, str(ML / "preparar_dados.py"), "--help"], cwd=self.temp.name, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("sinapi_es", result.stdout)
        self.assertIn("der_es", result.stdout)


if __name__ == "__main__":
    unittest.main()
