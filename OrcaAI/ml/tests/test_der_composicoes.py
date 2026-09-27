"""Casos de extração e reconciliação das CPUs DER, com valores conferidos à mão."""

import importlib
import io
import zipfile

import pandas as pd
import pytest
from openpyxl import Workbook


def _api():
    spec = importlib.util.find_spec("orca_ml.der_composicoes")
    assert spec is not None, "Leitor de composições DER ainda não implementado"
    return importlib.import_module("orca_ml.der_composicoes").ler_composicoes_der


def _line(**cells):
    row = [None] * 12
    for column, value in cells.items():
        row[ord(column) - ord("A")] = value
    return row


def _block(code, base=None, *, mo=(), mat=(), eq=(), child=(), total=0,
           totals=None, production=1, extra=()):
    totals = totals or {}
    rows = [
        _line(A="Serviço:\xa0", B=f"{code} - Serviço {code}", K="Unidade:\xa0", L="m2"),
        _line(A="Base:\xa0", B="LABOR - 2024", D="Código Base:", F=base or code,
              H="Fonte:", I="LABOR", K="Versão:", L=1),
    ]
    for section, items, letter in [
        ("Equipamento", eq, "A"), ("Mão-de-obra", mo, "B"),
        ("Itens de incidência", extra, "C"),
    ]:
        rows += [_line(A=section), *items, _line(A=f"({letter}) Total:", K=totals.get(letter, 0))]
    rows += [_line(A="(D) Produção da Equipe:", K=production),
             _line(A="(E) Custo Unitário da Execução [(A/D)+(B/D)+(C/D)]:", K=totals.get("E", 0))]
    for section, items, letter in [("Material", mat, "F"), ("Serviços", child, "G")]:
        rows += [_line(A=section), *items, _line(A=f"({letter}) Total:", K=totals.get(letter, 0))]
    rows += [_line(A="Itens de transporte"), _line(A="(H) Total:", K=totals.get("H", 0)),
             _line(A="Custo Direto Total (E)+(F)+(G)+(H):", K=total),
             _line(A="BDI:", K=0), _line(A="Preço Unitário Total:", K=total)]
    return rows


def _zip(tmp_path, blocks, resources, *, title="NÃO DESONERADOS", month="Maio/2026"):
    path = tmp_path / "der.zip"
    with zipfile.ZipFile(path, "w") as z:
        for name in ["composicoes", "insumos"]:
            wb = Workbook()
            ws = wb.active
            ws.title = "relatorio"
            ws["A6"] = f"TABELA DER-ES {title} (BDI=0%)"
            ws["A7"] = f"Data Base: {month}"
            if name == "composicoes":
                for row in blocks:
                    ws.append(row)
            else:
                for category, code, desc, unit, price in resources:
                    ws.append([f"Categoria: {category}"])
                    ws.append(["'" + code, desc, unit, price])
            stream = io.BytesIO()
            wb.save(stream)
            z.writestr(f"pasta/XLS/tab_DER-EDIFICAÇÕES_2026_05_{name}.xlsx", stream.getvalue())
    return path


def _services(*items):
    return pd.DataFrame([dict(codigo=code, custo=cost, unidade="m2", fonte="der_es",
                              competencia="2026-05", regime="sem_desoneracao")
                         for code, cost in items])


def test_expande_codigo_base_e_agrega_mo_sem_contar_subcomposicao_duas_vezes(tmp_path):
    # A referência 001 resolve para o item 990001. H/I de Serviços são coef/preço.
    painter = _line(A="PINTOR", D="010140", E="H", F=10, G=100, H=20, I=.3, K=6)
    material = _line(A="TINTA", F="037513", G="L", H=10, I=.25, K=2.5)
    blocks = _block("990001", "001", mo=[painter], total=6, totals={"B": 6, "E": 6})
    blocks += _block("190115", mo=[painter], mat=[material],
                     child=[_line(A="Auxiliar", F="'001", G="m2", H=.5, I=6, K=3)],
                     total=11.5, totals={"B": 6, "E": 6, "F": 2.5, "G": 3})
    path = _zip(tmp_path, blocks, [("Mão-de-obra", "010140", "PINTOR", "H", 10),
                                  ("Material", "037513", "TINTA", "L", 10)])
    services, resources, audit = _api()(path, "2026-05", _services(("190115", 11.5)))
    assert services.status_composicao.tolist() == ["ok"]
    assert services.custo_recalculado_hierarquia.tolist() == [11.5]
    assert services.diferenca_reconciliacao.tolist() == [0]
    assert len(resources) == 2
    mo = resources.set_index("codigo_recurso").loc["010140"]
    assert mo.coeficiente == pytest.approx(.45)
    assert mo.horas_mao_obra == pytest.approx(.45)
    assert mo.preco_efetivo == 20  # Encargos já incluídos; não usar salário-base 10.
    assert mo.custo_folha_sem_truncamento == pytest.approx(9)
    assert audit["servicos_ok"] == 1
    assert len(audit["arquivo_sha256"]) == 64


def test_trunca_no_nivel_correto_e_preserva_residuo_nas_folhas(tmp_path):
    blocks = _block("190115", mo=[_line(A="AJUDANTE", D="010101", E="H", F=9.13,
                    G=157.27, H=23.48, I=.3235, K=7.59)], total=7.59,
                    totals={"B": 7.59, "E": 7.59})
    path = _zip(tmp_path, blocks, [("Mão-de-obra", "010101", "AJUDANTE", "H", 9.13)])
    services, resources, _ = _api()(path, "2026-05", _services(("190115", 7.59)))
    assert services.status_composicao.tolist() == ["ok"]
    assert services.custo_recalculado_hierarquia.iloc[0] == 7.59
    assert resources.custo_folha_sem_truncamento.iloc[0] == pytest.approx(7.59578)


def test_equipamento_k_e_preco_ponderado_e_coeficiente_nao_sao_horas(tmp_path):
    blocks = _block("990237", eq=[_line(A="POLIDORA", D="081152", E="UND", F=.000064,
                    G=1, H=0, I=8410.08, J=0, K=8410.08)], total=.53,
                    totals={"A": .53, "E": .53})
    path = _zip(tmp_path, blocks, [("Equipamento", "081152", "POLIDORA", "UND", 8410.08)])
    services, resources, _ = _api()(path, "2026-05", _services(("990237", .53)))
    assert services.status_composicao.tolist() == ["ok"]
    r = resources.iloc[0]
    assert r.categoria == "equipamento"
    assert r.coeficiente == .000064
    assert r.custo_folha_sem_truncamento == pytest.approx(.53824512)
    assert r.horas_equipamento == 0


@pytest.mark.parametrize("problem", ["ciclo", "filho_ausente", "insumo_ausente", "divergencia", "preco_ausente"])
def test_exclui_composicao_incompleta_preservando_servico_e_motivo(tmp_path, problem):
    resource = _line(A="TINTA", F="037513", G="L", H=10, I=1, K=10)
    children = []
    if problem in ["ciclo", "filho_ausente"]:
        children = [_line(A="Auxiliar", F="190115" if problem == "ciclo" else "999999",
                          G="m2", H=1, I=10, K=10)]
    if problem == "preco_ausente":
        resource[7] = None
    blocks = _block("190115", mat=[] if children else [resource], child=children,
                     total=10, totals={"F": 0 if children else 10, "G": 10 if children else 0})
    path = _zip(tmp_path, blocks, [] if problem == "insumo_ausente" else
                [("Material", "037513", "TINTA", "L", 10)])
    s, r, a = _api()(path, "2026-05", _services(("190115", 11 if problem == "divergencia" else 10)))
    assert s.codigo.tolist() == ["190115"]
    assert s.status_composicao.tolist() == ["excluido"]
    assert s.motivo_exclusao.str.len().iloc[0] > 8
    assert r.empty
    assert a["servicos_excluidos"] == 1


def test_divisao_por_producao_so_afeta_equipe(tmp_path):
    blocks = _block("190115", mo=[_line(A="PINTOR", D="010140", E="H", F=10, G=100,
                    H=20, I=2, K=40)], mat=[_line(A="TINTA", F="037513", G="L", H=10, I=1, K=10)],
                    production=2, total=30, totals={"B": 40, "E": 20, "F": 10})
    path = _zip(tmp_path, blocks, [("Mão-de-obra", "010140", "PINTOR", "H", 10),
                                  ("Material", "037513", "TINTA", "L", 10)])
    s, r, _ = _api()(path, "2026-05", _services(("190115", 30)))
    assert s.status_composicao.tolist() == ["ok"]
    assert r.horas_mao_obra.sum() == 1
    assert r.custo_folha_sem_truncamento.sum() == 30


def test_rejeita_mes_regime_e_snapshot_misturados(tmp_path):
    path = _zip(tmp_path, _block("190115", total=0), [])
    with pytest.raises(ValueError, match="competência"):
        _api()(path, "2026-04", _services(("190115", 10)))
    mixed = pd.concat([_services(("190115", 10)), _services(("190115", 10))])
    with pytest.raises(ValueError, match="duplicad"):
        _api()(path, "2026-05", mixed)
    path = _zip(tmp_path, _block("190115", total=0), [], title="DESONERADOS")
    with pytest.raises(ValueError, match="regime"):
        _api()(path, "2026-05", _services(("190115", 10)))


def test_nao_aceita_mesmo_codigo_com_descricao_de_outro_servico(tmp_path):
    blocks = _block("190115", mat=[_line(A="TINTA", F="037513", G="L", H=10, I=1, K=10)],
                    total=10, totals={"F": 10})
    path = _zip(tmp_path, blocks, [("Material", "037513", "TINTA", "L", 10)])
    snapshot = _services(("190115", 10)).assign(descricao="Serviço diferente")
    s, r, _ = _api()(path, "2026-05", snapshot)
    assert s.status_composicao.tolist() == ["excluido"]
    assert "descrição" in s.motivo_exclusao.iloc[0].lower()
    assert r.empty


def test_snapshot_com_hash_de_outra_edicao_fica_excluido(tmp_path):
    blocks = _block("190115", mat=[_line(A="TINTA", F="037513", G="L", H=10, I=1, K=10)],
                    total=10, totals={"F": 10})
    path = _zip(tmp_path, blocks, [("Material", "037513", "TINTA", "L", 10)])
    snapshot = _services(("190115", 10)).assign(arquivo_sha256="f" * 64)
    s, r, _ = _api()(path, "2026-05", snapshot)
    assert s.status_composicao.tolist() == ["excluido"]
    assert "sha256" in s.motivo_exclusao.iloc[0].lower()
    assert r.empty
