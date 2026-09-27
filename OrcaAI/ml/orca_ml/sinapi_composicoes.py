"""Recursos SINAPI expandidos, com proveniência e conciliação de custos.

O total hierárquico trunca cada parcela a centavos. O custo de uma folha
expandida não aplica esse truncamento: serve para um perfil aproximado de
recursos, nunca para substituir o total publicado. Ausências não viram zero.
"""
from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from decimal import Decimal, ROUND_DOWN
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .sinapi import _cabecalho, _codigo, _codigo_csd, _decimal, _ler_analitico, _normalizado, _texto

CATEGORIAS = {
    'material': 'material', 'mao de obra': 'mao_obra',
    'encargos complementares': 'mao_obra',
    'equipamento (aquisicao)': 'equipamento',
    'equipamento (locacao)': 'equipamento',
    'servicos': 'outros', 'especiais': 'outros',
}
COLUNAS_RECURSOS = [
    'fonte', 'competencia', 'codigo_servico', 'codigo_recurso',
    'categoria_nativa', 'categoria', 'descricao_recurso', 'unidade_recurso',
    'coeficiente', 'preco_es', 'preco_sp', 'preco_efetivo', 'uf_preco',
    'custo_folha_sem_truncamento', 'horas_mao_obra', 'horas_equipamento',
    'localizador',
]


class ComposicaoInvalida(ValueError):
    """Uma composição incompleta não invalida outros serviços selecionados."""


def _numero_opcional(value, local):
    if _texto(value) in {'', '-', '–', '—'}:
        return None
    result = _decimal(value, local)
    if result < 0:
        raise ValueError(f'Número negativo em {local}')
    return result


def _insumos(sheet, competencia, member):
    iterator = sheet.iter_rows(values_only=True)
    header = [next(iterator) for _ in range(10)]
    ano, mes = competencia.split('-')
    if _texto(header[2][1]) != f'{mes}/{ano}':
        raise ValueError('Competência divergente em ISD')
    if 'sem desoneracao' not in _normalizado(header[1][0]):
        raise ValueError('Regime sem desoneração não confirmado em ISD')
    expected = ['classificacao', 'codigo do insumo', 'descricao do insumo', 'unidade', 'origem de preco']
    if [_normalizado(v) for v in header[9][:5]] != expected:
        raise ValueError('Cabeçalho inesperado em ISD')
    columns = {}
    for uf in ('ES', 'SP'):
        positions = [i for i, v in enumerate(header[9]) if _texto(v) == uf]
        if len(positions) != 1:
            raise ValueError(f'Coluna {uf} ausente ou duplicada em ISD')
        columns[uf] = positions[0]
    rows = {}
    for line, row in enumerate(iterator, 11):
        if all(v is None for v in row):
            continue
        code = _codigo(row[1], f'ISD!B{line}')
        if code in rows:
            raise ValueError(f'Insumo duplicado em ISD: {code}')
        if not _texto(row[2]) or not _texto(row[3]):
            raise ValueError(f'Descrição/unidade ausente em ISD!B{line}')
        record = {
            'categoria_nativa': _texto(row[0]), 'descricao_recurso': _texto(row[2]),
            'unidade_recurso': _texto(row[3]), 'linha': line,
        }
        for uf, col in columns.items():
            local = f'{member}!ISD!{get_column_letter(col + 1)}{line}'
            record[f'preco_{uf.lower()}'] = _numero_opcional(row[col], local)
            record[f'local_{uf.lower()}'] = local
        rows[code] = record
    return rows


def _custos(sheet, competencia, member):
    iterator, cost_col = _cabecalho(sheet, competencia, 'custos')
    rows = {}
    for line, row in enumerate(iterator, 11):
        if all(v is None for v in row):
            continue
        code = _codigo_csd(row[1], f'CSD!B{line}')
        if code in rows:
            raise ValueError(f'Composição duplicada em CSD: {code}')
        local = f'{member}!CSD!{get_column_letter(cost_col + 1)}{line}'
        fraction = _numero_opcional(row[cost_col + 1], local + ' %AS')
        if fraction is not None and fraction > 1:
            raise ValueError(f'%AS fora da faixa em {local}')
        rows[code] = {
            'grupo': _texto(row[0]), 'descricao': _texto(row[2]),
            'unidade': _texto(row[3]), 'custo': _numero_opcional(row[cost_col], local),
            'percentual_as': fraction, 'localizador_csd': local,
        }
    return rows


def _validar_selecao(servicos, competencia, sha):
    required = {'fonte', 'competencia', 'codigo', 'descricao', 'grupo', 'unidade', 'custo', 'regime'}
    if missing := required - set(servicos.columns):
        raise ValueError(f'Colunas ausentes no snapshot: {sorted(missing)}')
    result = servicos.copy(deep=True).reset_index(drop=True)
    for field, expected in [('fonte', 'sinapi_es'), ('competencia', competencia), ('regime', 'sem_desoneracao')]:
        if not result[field].eq(expected).all():
            raise ValueError(f'{field} divergente no snapshot SINAPI')
    result['codigo'] = [_codigo(v, 'snapshot.codigo') for v in result.codigo]
    if result.codigo.duplicated().any():
        raise ValueError('Serviço duplicado no snapshot')
    if 'arquivo_sha256' in result and not result.arquivo_sha256.eq(sha).all():
        raise ValueError('SHA-256 do snapshot diverge do ZIP selecionado')
    return result


def ler_composicoes_sinapi(path: str | Path, competencia: str, servicos: pd.DataFrame):
    """Retorne serviços preservados, recursos completos e auditoria JSON.

    Problemas localizados são registrados em ``status_composicao`` e não
    produzem recursos parciais. Erros de esquema, proveniência, duplicatas e
    números inválidos interrompem a leitura. O indicador %AS é conferido
    quanto à presença de contribuição SP; não se exige igualdade entre a
    fração oficial e a aproximação de folhas sem truncamento.
    """
    if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])', competencia):
        raise ValueError('Competência deve usar YYYY-MM')
    path = Path(path)
    raw = path.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    selected = _validar_selecao(servicos, competencia, sha)
    expected = f"SINAPI_Referência_{competencia.replace('-', '_')}.xlsx"
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if archive.testzip() is not None:
            raise ValueError('CRC inválido no ZIP SINAPI')
        members = [v for v in archive.namelist() if unicodedata.normalize('NFC', Path(v).name) == expected]
        if len(members) != 1:
            raise ValueError(f'Esperado um arquivo {expected}')
        member = members[0]
        wb = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=False)
        try:
            if not {'Analítico', 'ISD', 'CSD'}.issubset(wb.sheetnames):
                raise ValueError('Abas Analítico, ISD e CSD são obrigatórias')
            nodes = _ler_analitico(wb['Analítico'], competencia)
            edges = defaultdict(list)
            for line, row in enumerate(wb['Analítico'].iter_rows(values_only=True), 1):
                if line >= 11 and len(row) > 2 and _texto(row[2]):
                    code = _codigo(row[1], f'Analítico!B{line}')
                    edges[code].append(f'{member}!Analítico!G{line}')
            for code, node in nodes.items():
                for item, local in zip(node['itens'], edges[code], strict=True):
                    item['localizador'] = local
            insumos = _insumos(wb['ISD'], competencia, member)
            costs = _custos(wb['CSD'], competencia, member)
        finally:
            wb.close()

    selected['status_composicao'] = 'excluido'
    selected['motivo_exclusao'] = ''
    selected['custo_recalculado_hierarquia'] = float('nan')
    selected['diferenca_reconciliacao'] = float('nan')
    selected['localizador_analitico'] = ''
    resources, exclusions, approximations = [], [], []
    fallback_services = 0
    memo = {}

    def expand(code, active):
        if code in active:
            raise ComposicaoInvalida(f'ciclo: {code}')
        if code in memo:
            return memo[code]
        if code not in nodes:
            raise ComposicaoInvalida(f'auxiliar_ausente: {code}')
        node = nodes[code]
        if not node['itens']:
            raise ComposicaoInvalida(f'composicao_sem_itens: {code}')
        leaves = defaultdict(Decimal)
        paths = defaultdict(set)
        total = Decimal(0)
        active = active | {code}
        for edge in node['itens']:
            qty = Decimal(edge['coeficiente'])
            if qty == 0:
                continue
            item = edge['codigo']
            if edge['tipo'] == 'composicao':
                if item in active:
                    raise ComposicaoInvalida(f'ciclo: {item}')
                if item not in nodes:
                    raise ComposicaoInvalida(f'auxiliar_ausente: {item}')
                if _normalizado(nodes[item]['unidade']) != edge['unidade']:
                    raise ComposicaoInvalida(f'unidade_divergente: auxiliar {item}')
                child_cost, child_leaves, child_paths = expand(item, active)
                for resource, coefficient in child_leaves.items():
                    leaves[resource] += qty * coefficient
                    paths[resource].update(child_paths[resource])
                    paths[resource].add(edge['localizador'])
                total += (qty * child_cost).quantize(Decimal('.01'), rounding=ROUND_DOWN)
            else:
                if item not in insumos:
                    raise ComposicaoInvalida(f'insumo_ausente: {item}')
                ins = insumos[item]
                if _normalizado(ins['categoria_nativa']) not in CATEGORIAS:
                    raise ComposicaoInvalida(f'categoria_desconhecida: {item}: {ins["categoria_nativa"]}')
                if _normalizado(ins['unidade_recurso']) != edge['unidade']:
                    raise ComposicaoInvalida(f'unidade_divergente: insumo {item}')
                price = ins['preco_es'] if ins['preco_es'] is not None else ins['preco_sp']
                if price is None:
                    raise ComposicaoInvalida(f'sem_preco: {item}')
                leaves[item] += qty
                paths[item].add(edge['localizador'])
                total += (qty * price).quantize(Decimal('.01'), rounding=ROUND_DOWN)
        if not leaves:
            raise ComposicaoInvalida(f'composicao_sem_recursos_positivos: {code}')
        memo[code] = total, dict(leaves), dict(paths)
        return memo[code]

    for idx, service in selected.iterrows():
        code = service.codigo
        try:
            if code not in costs or code not in nodes:
                raise ComposicaoInvalida('composicao_ausente')
            official = costs[code]
            if official['custo'] is None or official['custo'] <= 0:
                raise ComposicaoInvalida('custo_oficial_ausente_ou_nao_positivo')
            for field in ('descricao', 'unidade', 'grupo'):
                if _normalizado(service[field]) != _normalizado(official[field]):
                    raise ComposicaoInvalida(f'snapshot_divergente: {field}')
            if _decimal(service.custo, 'snapshot.custo') != official['custo']:
                raise ComposicaoInvalida('snapshot_divergente: custo')
            if _normalizado(nodes[code]['descricao']) != _normalizado(official['descricao']) or _normalizado(nodes[code]['unidade']) != _normalizado(official['unidade']):
                raise ComposicaoInvalida('analitico_divergente_do_csd')
            if 'percentual_as' in service and pd.notna(service.percentual_as):
                if official['percentual_as'] is None or abs(_decimal(service.percentual_as, 'snapshot.%AS') - official['percentual_as']) > Decimal('0.000000000001'):
                    raise ComposicaoInvalida('snapshot_divergente: percentual_as')
            selected.at[idx, 'localizador_analitico'] = f"{member}!Analítico!B{nodes[code]['linha']}"
            calculated, leaves, paths = expand(code, set())
            selected.at[idx, 'custo_recalculado_hierarquia'] = float(calculated)
            difference = calculated - official['custo']
            selected.at[idx, 'diferenca_reconciliacao'] = float(difference)
            if difference != 0:
                raise ComposicaoInvalida(f'custo_nao_conciliado: diferenca={difference}')
            uses_sp = any(insumos[k]['preco_es'] is None and insumos[k]['preco_sp'] > 0 for k in leaves)
            if official['percentual_as'] is None or (official['percentual_as'] > 0) != uses_sp:
                raise ComposicaoInvalida('percentual_as_inconsistente_com_origem_dos_precos')
            total_leaf = Decimal(0)
            for item, qty in sorted(leaves.items()):
                ins = insumos[item]
                uf = 'ES' if ins['preco_es'] is not None else 'SP'
                price = ins[f'preco_{uf.lower()}']
                category = CATEGORIAS[_normalizado(ins['categoria_nativa'])]
                unit = _normalizado(ins['unidade_recurso'])
                leaf_cost = qty * price
                total_leaf += leaf_cost
                resources.append({
                    'fonte': 'sinapi_es', 'competencia': competencia,
                    'codigo_servico': code, 'codigo_recurso': item,
                    'categoria_nativa': ins['categoria_nativa'], 'categoria': category,
                    'descricao_recurso': ins['descricao_recurso'], 'unidade_recurso': ins['unidade_recurso'],
                    'coeficiente': float(qty),
                    'preco_es': None if ins['preco_es'] is None else float(ins['preco_es']),
                    'preco_sp': None if ins['preco_sp'] is None else float(ins['preco_sp']),
                    'preco_efetivo': float(price), 'uf_preco': uf,
                    'custo_folha_sem_truncamento': float(leaf_cost),
                    'horas_mao_obra': float(qty) if _normalizado(ins['categoria_nativa']) == 'mao de obra' and unit == 'h' else 0.,
                    'horas_equipamento': float(qty) if category == 'equipamento' and unit == 'h' else 0.,
                    'localizador': '; '.join([ins[f'local_{uf.lower()}']] + sorted(paths[item])),
                })
            approximations.append(float((total_leaf - calculated) / calculated * 100))
            fallback_services += int(uses_sp)
            selected.at[idx, 'status_composicao'] = 'ok'
        except ComposicaoInvalida as exc:
            reason = str(exc)
            selected.at[idx, 'motivo_exclusao'] = reason
            exclusions.append({'codigo': code, 'motivo': reason})

    resource_frame = pd.DataFrame(resources, columns=COLUNAS_RECURSOS)
    audit = {
        'fonte': 'sinapi_es', 'competencia': competencia, 'regime': 'sem_desoneracao',
        'arquivo': str(path.resolve()), 'arquivo_sha256': sha, 'arquivo_bytes': len(raw),
        'membro_xlsx': member, 'zip_crc_verificado': True,
        'composicoes_analitico': len(nodes), 'insumos_isd': len(insumos),
        'servicos_selecionados': len(selected), 'servicos_validos': len(selected) - len(exclusions),
        'servicos_excluidos': len(exclusions), 'recursos_expandidos': len(resources),
        'servicos_com_fallback_sp': fallback_services, 'exclusoes': exclusions,
        'excluidos_por_motivo': dict(Counter(v['motivo'].split(':')[0] for v in exclusions)),
        'diferenca_folhas_pct_min': min(approximations) if approximations else None,
        'diferenca_folhas_pct_max': max(approximations) if approximations else None,
        'convencoes': {
            'reconciliacao': 'Decimal; cada coeficiente × preço truncado em duas casas (ROUND_DOWN), em cada nível; diferença exigida zero.',
            'perfil_monetario': 'Soma de folhas sem truncamento. Participações derivadas são aproximadas; não há rateio para forçar igualdade ao custo CSD.',
            'fallback_sp': 'Somente se preço ES ausente e SP disponível; preço ES continua ausente. Indicador oficial %AS conferido quanto à existência de contribuição positiva SP, sem equiparar sua fração ao perfil aproximado.',
            'mao_obra': 'MAO DE OBRA e ENCARGOS COMPLEMENTARES; horas físicas somente da primeira categoria com unidade H.',
            'horas_equipamento': 'Coeficientes de folhas EQUIPAMENTO com unidade H; aquisições em UN não são convertidas em horas.',
            'exclusoes': 'Todo serviço selecionado é preservado. Serviço excluído não produz linhas de recursos parciais.',
        },
    }
    return selected, resource_frame, audit
