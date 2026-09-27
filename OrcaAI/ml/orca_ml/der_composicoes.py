"""CPUs DER-ES: expansão auditável dos recursos e custos de uma edição fixa.

Os coeficientes das folhas preservam os quantitativos físicos. Os custos das
folhas não são rateados para forçar o total publicado: diferenças decorrentes
de truncamento ficam explícitas. A reconciliação é feita separadamente, em
cada nível da composição, com os totais publicados do relatório.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation, ROUND_DOWN
import hashlib
import io
from pathlib import Path
import re
import unicodedata
import zipfile

import pandas as pd
from openpyxl import load_workbook

from .der import MESES


_ZERO = Decimal("0")
_CENT = Decimal("0.01")
_CATEGORIES = {
    "Mão-de-obra": "mao_obra", "Material": "material",
    "Equipamento": "equipamento", "Serv.Terceiro": "outros",
}
_SECTIONS = {
    "Equipamento": "A", "Mão-de-obra": "B", "Itens de incidência": "C",
    "Material": "F", "Serviços": "G", "Itens de transporte": "H",
}
_RESOURCE_COLUMNS = [
    "fonte", "competencia", "codigo_servico", "codigo_recurso",
    "categoria_nativa", "categoria", "descricao_recurso", "unidade_recurso",
    "coeficiente", "preco_efetivo", "uf_preco", "custo_folha_sem_truncamento",
    "horas_mao_obra", "horas_equipamento", "localizador",
]


def _text(value) -> str:
    return unicodedata.normalize("NFC", str(value if value is not None else "")).replace("\xa0", " ").strip()


def _code(value) -> str:
    return _text(value).lstrip("'")


def _unit(value) -> str:
    return _text(value).upper().replace("²", "2").replace("³", "3")


def _num(value, where: str, *, positive=False) -> Decimal:
    text = _text(value).replace(" ", "").replace("R$", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Preço/coeficiente ausente ou inválido em {where}: {value!r}") from exc
    if not number.is_finite() or number < 0 or (positive and number <= 0):
        raise ValueError(f"Valor não finito, negativo ou nulo em {where}: {value!r}")
    return number


def _trunc(value: Decimal) -> Decimal:
    return value.quantize(_CENT, rounding=ROUND_DOWN)


def _check_equal(actual: Decimal, expected, where: str) -> None:
    if actual != _num(expected, where):
        raise ValueError(f"Divergência de reconciliação em {where}: calculado={actual}, publicado={expected}")


def _validate_header(rows, competence: str, member: str) -> None:
    header = " ".join(_text(v) for row in rows[:7] for v in row)
    normal = unicodedata.normalize("NFKD", header).encode("ascii", "ignore").decode().lower()
    if not re.search(r"\bnao\s+desonerados\b", normal):
        raise ValueError(f"regime sem desoneração não confirmado em {member}")
    dates = re.findall(r"\b([a-z]+)\s*/\s*(20\d\d)\b", normal)
    year, month = map(int, competence.split("-"))
    if not any(MESES.get(m) == month and int(y) == year for m, y in dates):
        raise ValueError(f"competência {competence} não confirmada em {member}")
    bdis = re.findall(r"bdi\s*[:=]\s*([\d,.]+)", normal)
    if not bdis or any(_num(v, "BDI") != 0 for v in bdis):
        raise ValueError(f"BDI zero não confirmado em {member}")


def _read_rows(archive, suffix: str, competence: str):
    names = [n for n in archive.namelist() if re.search(rf"/XLS/[^/]*_{suffix}\.xlsx$", n, re.I)]
    if len(names) != 1:
        raise ValueError(f"Esperado um XLSX de {suffix}; encontrados {len(names)}")
    member = names[0]
    payload = archive.read(member)
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"Esperada uma aba em {member}")
        sheet = workbook.active
        rows = [tuple(row) + (None,) * max(0, 12 - len(row)) for row in sheet.values]
        _validate_header(rows, competence, member)
        return rows, member, sheet.title, hashlib.sha256(payload).hexdigest()
    finally:
        workbook.close()


def _parse_resources(rows, member):
    resources = {}
    category = None
    for number, row in enumerate(rows, 1):
        text = _text(row[0])
        if text.startswith("Categoria:"):
            category = text.split(":", 1)[1].strip()
        elif _code(row[0]).isdigit() and row[2]:
            code = _code(row[0])
            if code in resources:
                raise ValueError(f"Insumo duplicado {code} em {member}")
            resources[code] = {
                "categoria_nativa": category, "categoria": _CATEGORIES.get(category),
                "descricao": _text(row[1]), "unidade": _unit(row[2]),
                "preco": row[3], "localizador": f"{member}!A{number}",
            }
    return resources


def _parse_blocks(rows, member):
    blocks = {}
    bases = defaultdict(list)
    block, section = None, None
    for number, row in enumerate(rows, 1):
        text = _text(row[0])
        location = f"{member}!A{number}"
        if text == "Serviço:":
            code, description = _text(row[1]).split(" - ", 1)
            if code in blocks:
                raise ValueError(f"Serviço duplicado {code} em {location}")
            block = dict(codigo=code, descricao=description, unidade=_unit(row[11]),
                         localizador=location, items=[], totals={}, issues=[])
            blocks[code] = block
            section = None
        elif block is None:
            continue
        elif text == "Base:":
            base = _code(row[5])
            block.update(codigo_base=base, base=_text(row[1]), fonte=_text(row[8]))
            bases[base].append(block["codigo"])
        elif text in _SECTIONS:
            section = _SECTIONS[text]
        elif re.match(r"^\([A-H]\)", text):
            block["totals"][text[1]] = row[10]
        elif text.startswith("Custo Direto Total"):
            block["totals"]["direto"] = row[10]
        elif text == "BDI:":
            block["totals"]["bdi"] = row[10]
        elif text == "Preço Unitário Total:":
            block["totals"]["preco"] = row[10]
        elif text.startswith(("Planilha:", "Custo Horário da Execução")) or not text:
            continue
        elif section:
            code_col = 3 if section in {"A", "B", "C", "H"} else 5
            if _code(row[code_col]).isdigit():
                block["items"].append((section, row, location))
            else:
                block["issues"].append(f"Linha de recurso não reconhecida em {location}")
    return blocks, bases


def _expand(code, blocks, bases, resources, ancestry=()):
    if code in ancestry:
        raise ValueError(f"Ciclo de subcomposições: {' > '.join(ancestry + (code,))}")
    if code not in blocks:
        raise ValueError(f"Composição ausente: {code}")
    block = blocks[code]
    where = block["localizador"]
    if block["issues"]:
        raise ValueError("; ".join(block["issues"]))
    if block.get("fonte") != "LABOR" or not block.get("codigo_base"):
        raise ValueError(f"Código Base LABOR não confirmado em {where}")
    totals = block["totals"]
    production = _num(totals.get("D"), f"produção {where}", positive=True)
    _check_equal(_ZERO, totals.get("bdi"), f"BDI {where}")
    sums = {letter: _ZERO for letter in ("A", "B", "C", "F", "G", "H")}
    leaves = []
    for section, row, location in block["items"]:
        if section in {"C", "H"}:
            raise ValueError(f"Incidência/transporte ainda não suportado em {location}")
        if section == "G":
            child_base = _code(row[5])
            matches = bases.get(child_base, [])
            if len(matches) != 1:
                raise ValueError(f"Subcomposição ausente ou ambígua {child_base} em {location}")
            child_code = matches[0]
            if _unit(row[6]) != blocks[child_code]["unidade"]:
                raise ValueError(f"Unidade de subcomposição divergente em {location}")
            coefficient = _num(row[7], f"coeficiente {location}")
            child_total, child_leaves = _expand(child_code, blocks, bases, resources, ancestry + (code,))
            _check_equal(child_total, row[8], f"preço da subcomposição {location}")
            cost = _trunc(coefficient * child_total)
            _check_equal(cost, row[10], f"subtotal {location}")
            sums[section] += cost
            for leaf in child_leaves:
                leaves.append({**leaf, "coeficiente": leaf["coeficiente"] * coefficient,
                               "localizador": f"{location} > {leaf['localizador']}"})
            continue

        resource_code = _code(row[3] if section in {"A", "B"} else row[5])
        if resource_code not in resources:
            raise ValueError(f"Insumo ausente no catálogo: {resource_code} em {location}")
        resource = resources[resource_code]
        if resource["categoria"] is None:
            raise ValueError(f"Categoria de recurso desconhecida em {resource['localizador']}")
        unit = _unit(row[4] if section in {"A", "B"} else row[6])
        if unit != resource["unidade"]:
            raise ValueError(f"Unidade do insumo divergente em {location}")
        _num(resource["preco"], resource["localizador"])
        if section == "A":
            coefficient = _num(row[5], f"quantidade {location}")
            price = (_num(row[6], location) * _num(row[8], location)
                     + _num(row[7], location) * _num(row[9], location))
            _check_equal(price, row[10], f"preço ponderado do equipamento {location}")
            cost = _trunc(coefficient * price)
        else:
            coefficient = _num(row[8], f"consumo {location}")
            price = _num(row[7], f"preço efetivo {location}")
            cost = _trunc(coefficient * price)
            _check_equal(cost, row[10], f"subtotal {location}")
        expected = {"A": "equipamento", "B": "mao_obra"}.get(section)
        if expected and resource["categoria"] != expected:
            raise ValueError(f"Categoria da CPU diverge do catálogo em {location}")
        sums[section] += cost
        leaves.append({"codigo_recurso": resource_code, **resource,
                       "coeficiente": coefficient / production if section in {"A", "B"} else coefficient,
                       "preco_efetivo": price, "localizador": location})

    for section, calculated in sums.items():
        _check_equal(calculated, totals.get(section), f"total ({section}) {where}")
    execution = _trunc((sums["A"] + sums["B"] + sums["C"]) / production)
    _check_equal(execution, totals.get("E"), f"total (E) {where}")
    total = execution + sums["F"] + sums["G"] + sums["H"]
    _check_equal(total, totals.get("direto"), f"custo direto {where}")
    _check_equal(total, totals.get("preco"), f"preço total {where}")
    return total, leaves


def _aggregate(leaves, service_code, competence):
    groups = defaultdict(list)
    for leaf in leaves:
        groups[leaf["codigo_recurso"]].append(leaf)
    result = []
    for code, group in sorted(groups.items()):
        first = group[0]
        if len({(r["categoria"], r["unidade"]) for r in group}) != 1:
            raise ValueError(f"Recurso com unidade/categoria inconsistente: {code}")
        coefficient = sum((r["coeficiente"] for r in group), _ZERO)
        cost = sum((r["coeficiente"] * r["preco_efetivo"] for r in group), _ZERO)
        # Zero de consumo é ausência, sem inventar preço efetivo por divisão.
        price = cost / coefficient if coefficient else first["preco_efetivo"]
        hours = coefficient if first["unidade"] == "H" else _ZERO
        result.append({
            "fonte": "der_es", "competencia": competence, "codigo_servico": service_code,
            "codigo_recurso": code, "categoria_nativa": first["categoria_nativa"],
            "categoria": first["categoria"], "descricao_recurso": first["descricao"],
            "unidade_recurso": first["unidade"], "coeficiente": float(coefficient),
            "preco_efetivo": float(price), "uf_preco": "ES",
            "custo_folha_sem_truncamento": float(cost),
            "horas_mao_obra": float(hours if first["categoria"] == "mao_obra" else _ZERO),
            "horas_equipamento": float(hours if first["categoria"] == "equipamento" else _ZERO),
            "localizador": " | ".join(sorted({r["localizador"] for r in group})),
        })
    return result


def ler_composicoes_der(path, competencia: str, servicos: pd.DataFrame):
    """Leia e confira recursos-folha dos serviços selecionados de uma edição.

    Retorna ``(servicos_df, recursos_df, auditoria_dict)``. Falhas de uma CPU
    excluem apenas aquele serviço do experimento e preservam sua linha com a
    justificativa. Regime/edição incoerentes e entradas duplicadas rejeitam a
    chamada inteira. O snapshot precisa ter ``codigo``, ``unidade`` e ``custo``.
    """
    if not re.fullmatch(r"20\d\d-(?:0[1-9]|1[0-2])", competencia):
        raise ValueError("competência deve ser YYYY-MM")
    if not {"codigo", "unidade", "custo"}.issubset(servicos.columns):
        raise ValueError("Snapshot requer codigo, unidade e custo")
    selected = servicos.copy().reset_index(drop=True)
    selected["codigo"] = selected["codigo"].map(_code)
    if selected["codigo"].duplicated().any():
        raise ValueError("Serviço duplicado no snapshot")
    for column, expected in [("competencia", competencia), ("fonte", "der_es"), ("regime", "sem_desoneracao")]:
        if column in selected and not selected[column].eq(expected).all():
            raise ValueError(f"Snapshot contém {column}/competência incompatível: esperado {expected}")
    path = Path(path)
    archive_sha = hashlib.sha256(path.read_bytes()).hexdigest()
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("ZIP DER corrompido")
        cpu_rows, cpu_member, cpu_sheet, cpu_sha = _read_rows(archive, "composicoes", competencia)
        resource_rows, resource_member, resource_sheet, resource_sha = _read_rows(archive, "insumos", competencia)
    resources = _parse_resources(resource_rows, resource_member)
    blocks, bases = _parse_blocks(cpu_rows, cpu_member)
    records, all_resources = [], []
    for service in selected.to_dict("records"):
        code = service["codigo"]
        out = {**service, "status_composicao": "excluido", "motivo_exclusao": "",
               "custo_recalculado_hierarquia": float("nan"), "diferenca_reconciliacao": float("nan")}
        try:
            if code not in blocks:
                raise ValueError(f"Composição ausente: {code}")
            if "arquivo_sha256" in service and _text(service["arquivo_sha256"]) != archive_sha:
                raise ValueError(f"SHA256 do snapshot diverge do ZIP de composições: {code}")
            if "descricao" in service and _text(service["descricao"]) != blocks[code]["descricao"]:
                raise ValueError(f"Descrição do serviço diverge da CPU: {code}")
            if _unit(service["unidade"]) != blocks[code]["unidade"]:
                raise ValueError(f"Unidade do serviço diverge da CPU: {code}")
            total, leaves = _expand(code, blocks, bases, resources)
            out["custo_recalculado_hierarquia"] = float(total)
            difference = total - _num(service["custo"], f"custo do snapshot {code}", positive=True)
            out["diferenca_reconciliacao"] = float(difference)
            _check_equal(total, service["custo"], f"CPU versus snapshot {code}")
            if not leaves or sum((x["coeficiente"] * x["preco_efetivo"] for x in leaves), _ZERO) <= 0:
                raise ValueError(f"Composição sem recursos com custo positivo: {code}")
            flat = _aggregate(leaves, code, competencia)
            all_resources.extend(flat)
            out["status_composicao"] = "ok"
        except ValueError as exc:
            out["motivo_exclusao"] = str(exc)
        records.append(out)
    services_df = pd.DataFrame(records, columns=list(selected.columns) + [
        "status_composicao", "motivo_exclusao", "custo_recalculado_hierarquia", "diferenca_reconciliacao"])
    resources_df = pd.DataFrame(all_resources, columns=_RESOURCE_COLUMNS)
    counts = Counter(services_df.status_composicao)
    audit = {
        "fonte": "der_es", "competencia": competencia, "uf": "ES", "regime": "sem_desoneracao",
        "arquivo": str(path.resolve()), "arquivo_sha256": archive_sha,
        "planilhas": [{"membro": cpu_member, "aba": cpu_sheet, "sha256": cpu_sha},
                      {"membro": resource_member, "aba": resource_sheet, "sha256": resource_sha}],
        "composicoes_no_catalogo": len(blocks), "insumos_no_catalogo": len(resources),
        "servicos_selecionados": len(selected), "servicos_ok": counts["ok"],
        "servicos_excluidos": counts["excluido"], "linhas_recursos_folha": len(resources_df),
        "recursos_distintos": int(resources_df.codigo_recurso.nunique()),
        "tolerancia_reconciliacao_reais": 0,
        "convencoes": [
            "Subcomposição resolve pelo Código Base LABOR, não pelo item do orçamento.",
            "Na seção Serviços, H=coeficiente e I=custo, apesar dos cabeçalhos invertidos.",
            "Equipamento: F=quantidade; K=preço ponderado, conferido com G*I + H*J.",
            "Mão de obra usa preço com encargos da coluna H, sem aplicar encargos novamente.",
            "ROUND_DOWN a centavos por item e em cada nível; totais de A..H, CPU e snapshot conferidos.",
            "Coeficientes de equipe divididos pela produção D; material e serviços não são divididos por D.",
            "Custos de recursos-folha preservam coeficiente*preço sem truncamento; suas participações são aproximadas e devem ser normalizadas pela própria soma, sem rateio do residual.",
            "Horas somadas somente para recursos cuja unidade é H, dentro da categoria correspondente.",
            "uf_preco=ES identifica a tabela estadual; não comprova coleta local individual de cada insumo.",
        ],
        "exclusoes": [{"codigo": r["codigo"], "motivo": r["motivo_exclusao"]}
                      for r in records if r["status_composicao"] != "ok"],
    }
    return services_df, resources_df, audit
