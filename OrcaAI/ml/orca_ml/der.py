"""Importação dos preços de serviços de edificações DER-ES, sem desoneração.

Os ZIPs originais são preservados. A assinatura disponível é apenas de
descrição e unidade: ela não demonstra estabilidade da composição analítica.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import subprocess
import tempfile
import unicodedata
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

from openpyxl import load_workbook


PAGINA = "https://der.es.gov.br/referencial-de-precos-edificacoes"
COLUNAS = [
    "fonte", "competencia", "codigo", "descricao", "unidade", "grupo", "custo",
    "regime", "publicado_em", "arquivo_sha256", "localizador", "assinatura_tecnica",
    "assinatura_tipo",
]
MESES = {
    "janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4,
    "maio": 5, "junho": 6, "julho": 7, "agosto": 8, "setembro": 9,
    "outubro": 10, "novembro": 11, "dezembro": 12,
}


class _Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hrefs: list[str] = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.hrefs.append(href)


def extrair_links_oficiais(html: str) -> dict[str, str]:
    """Retorne competências e URLs observadas no índice oficial, sem inferir nomes."""
    parser = _Links()
    parser.feed(html)
    links = {}
    for href in parser.hrefs:
        url = urljoin(PAGINA, href)
        parsed = urlparse(url)
        path = unquote(parsed.path)
        if parsed.scheme != "https" or parsed.netloc != "der.es.gov.br":
            continue
        if "Referencial de Edificações/OBRAS_" not in path or not path.lower().endswith(".zip"):
            continue
        found = re.search(r"TABELA REFERENCIAL OBRAS - DER-ES - (20\d\d-(?:0[1-9]|1[0-2])) - ", path, re.I)
        if found:
            competencia = found.group(1)
            if competencia in links and links[competencia] != url:
                raise ValueError(f"Múltiplas edições para {competencia}; seleção manual necessária")
            links[competencia] = url
    return dict(sorted(links.items()))


def _texto(value) -> str:
    return unicodedata.normalize("NFC", str(value or "")).replace("\xa0", " ").strip()


def _numero(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = _texto(value).replace("R$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    return float(text)


def _validar_cabecalho(rows: list[tuple], competencia: str) -> None:
    cabecalho = " ".join(_texto(v) for row in rows[:12] for v in row)
    norm = unicodedata.normalize("NFKD", cabecalho).encode("ascii", "ignore").decode().lower()
    if not re.search(r"\bnao\s+desonerados\b", norm):
        raise ValueError("regime sem desoneração não confirmado no relatório")
    ano, mes = map(int, competencia.split("-"))
    bases = re.findall(r"data\s+base\s*:\s*([a-zçãõéê]+)\s*/\s*(20\d\d)", norm)
    if not bases or all((MESES.get(m), int(y)) != (mes, ano) for m, y in bases):
        raise ValueError(f"competência {competencia} diverge da Data Base")
    bdis = re.findall(r"bdi\s*[:=]\s*([\d,.]+)\s*%?", norm)
    if not bdis or any(_numero(bdi) != 0 for bdi in bdis):
        raise ValueError("BDI zero não confirmado no relatório")
    header = [_texto(v).lower() for v in rows[11][:7]]
    expected = ["item", "fonte/código", "especificação do serviço", "und.", "quant.", "preço unitário", "preço total"]
    if header != expected:
        raise ValueError(f"Cabeçalho de serviços inesperado: {header}")


def parse_der_zip(arquivo: Path, competencia: str, publicado_em: str = "") -> list[dict]:
    """Leia serviços no ZIP, verifique regime/data/BDI e preserve localização."""
    arquivo = Path(arquivo)
    if not re.fullmatch(r"20\d\d-(?:0[1-9]|1[0-2])", competencia):
        raise ValueError("competência deve ser YYYY-MM")
    if publicado_em and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", publicado_em):
        raise ValueError("publicado_em deve ser ISO date")
    sha = hashlib.sha256(arquivo.read_bytes()).hexdigest()
    with zipfile.ZipFile(arquivo) as archive:
        if archive.testzip() is not None:
            raise ValueError(f"ZIP inválido: {arquivo}")
        names = [n for n in archive.namelist() if re.search(r"/XLS/[^/]*_servicos\.xlsx$", n, re.I)]
        if len(names) != 1:
            raise ValueError(f"Esperado um relatório XLSX de serviços; encontrados {len(names)}")
        member = names[0]
        workbook = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
        try:
            if len(workbook.sheetnames) != 1:
                raise ValueError("Quantidade inesperada de abas no relatório de serviços")
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            first = [next(iterator) for _ in range(12)]
            _validar_cabecalho(first, competencia)
            grupos: list[tuple[str, str]] = []
            rows = []
            for row_no, row in enumerate(iterator, start=13):
                if len(row) < 7:
                    continue
                codigo = _texto(row[0]).lstrip("'")
                descricao = _texto(row[2])
                unidade = _texto(row[3])
                preco = row[5]
                if not codigo or not descricao or not re.fullmatch(r"\d+", codigo):
                    continue
                if not unidade and preco is None:
                    while grupos and not codigo.startswith(grupos[-1][0]):
                        grupos.pop()
                    if grupos and grupos[-1][0] == codigo:
                        grupos.pop()
                    grupos.append((codigo, descricao))
                    continue
                if not unidade or preco is None:
                    raise ValueError(f"Serviço incompleto em {member}!A{row_no}")
                custo = _numero(preco)
                if custo <= 0:
                    raise ValueError(f"Custo não positivo em {member}!F{row_no}")
                assinatura_base = "|".join(
                    unicodedata.normalize("NFC", value).casefold().strip()
                    for value in (descricao, unidade)
                )
                rows.append({
                    "fonte": "der_es", "competencia": competencia, "codigo": codigo,
                    "descricao": descricao, "unidade": unidade,
                    "grupo": " / ".join(g[1] for g in grupos), "custo": custo,
                    "regime": "sem_desoneracao", "publicado_em": publicado_em,
                    "arquivo_sha256": sha, "localizador": f"{member}!A{row_no}",
                    "assinatura_tecnica": hashlib.sha256(assinatura_base.encode()).hexdigest(),
                    "assinatura_tipo": "descricao_unidade",
                })
        finally:
            workbook.close()
    if not rows:
        raise ValueError(f"Nenhum serviço extraído de {arquivo}")
    return rows


def preparar_csv(rows: list[dict], destino: Path) -> None:
    """Grave a série canônica em UTF-8 após verificar o esquema e unicidade."""
    if not rows:
        raise ValueError("Nenhum serviço para preparar")
    keys = [(r["fonte"], r["competencia"], r["codigo"], r["regime"]) for r in rows]
    if len(keys) != len(set(keys)):
        raise ValueError("Código/competência duplicado; revisão exige seleção explícita")
    if any(set(r) != set(COLUNAS) for r in rows):
        raise ValueError("Esquema canônico incompleto ou com colunas extras")
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", newline="", dir=destino.parent, delete=False) as stream:
        tmp = Path(stream.name)
        writer = csv.DictWriter(stream, fieldnames=COLUNAS)
        writer.writeheader()
        writer.writerows(sorted(rows, key=lambda r: (r["competencia"], r["codigo"])))
    tmp.replace(destino)


def _curl(url: str, output: Path | None = None) -> str | None:
    command = ["curl", "--fail", "--location", "--silent", "--show-error", "--retry", "2", "--max-time", "180", "--proto", "=https", "--proto-redir", "=https"]
    if output is not None:
        subprocess.run(command + ["--output", str(output), url], check=True)
        return None
    return subprocess.run(command + [url], check=True, capture_output=True).stdout.decode("utf-8")


def coletar_historico(
    base_dir: Path, csv_destino: Path, inicio: str = "2025-01", fim: str | None = None,
) -> dict[str, int]:
    """Baixe as edições observadas no portal e gere a série de serviços."""
    links = extrair_links_oficiais(_curl(PAGINA))
    selected = {c: u for c, u in links.items() if c >= inicio and (fim is None or c <= fim)}
    if not selected:
        raise ValueError("Nenhuma competência no intervalo no índice oficial")
    all_rows = []
    counts = {}
    for competencia, url in selected.items():
        name = Path(unquote(urlparse(url).path)).name
        directory = Path(base_dir) / competencia / "originais"
        directory.mkdir(parents=True, exist_ok=True)
        target = directory / name
        if not target.exists():
            partial = directory / (name + ".part")
            try:
                _curl(url, partial)
                with zipfile.ZipFile(partial) as archive:
                    if archive.testzip() is not None:
                        raise ValueError(f"ZIP corrompido: {competencia}")
                partial.replace(target)
            finally:
                partial.unlink(missing_ok=True)
        rows = parse_der_zip(target, competencia)
        counts[competencia] = len(rows)
        all_rows.extend(rows)
        manifest = {
            "fonte": "DER-ES", "competencia": competencia, "regime": "sem_desoneracao",
            "url_oficial_observada": url, "pagina_indice": PAGINA,
            "capturado_em_utc": datetime.now(timezone.utc).isoformat(),
            "arquivo": name, "bytes": target.stat().st_size,
            "sha256": rows[0]["arquivo_sha256"], "zip_crc_verificado": True,
            "servicos_extraidos": len(rows), "bdi_planilha": 0,
            "publicado_em": None,
            "assinatura_tipo": "descricao_unidade",
            "limite_assinatura": "Não verifica insumos, coeficientes nem subcomposições da CPU.",
        }
        (directory / "manifesto.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"{competencia}: {len(rows)} serviços, SHA-256 {manifest['sha256']}", flush=True)
    preparar_csv(all_rows, csv_destino)
    return counts
