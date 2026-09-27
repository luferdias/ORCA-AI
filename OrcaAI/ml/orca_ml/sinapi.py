"""Leia SINAPI nacional em memória e prepare custos ES sem desoneração.

O código nas células HYPERLINK é validado contra a aba Analítico, sem usar
o cache de fórmulas. Assinaturas incluem a hierarquia técnica, mas não
preços, situação de preços ou grupo classificatório. Os ZIPs ficam intactos.
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import re
import tempfile
import unicodedata
import zipfile
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


COLUNAS = [
    "fonte", "competencia", "codigo", "descricao", "unidade", "grupo", "custo",
    "regime", "publicado_em", "arquivo_sha256", "localizador", "assinatura_tecnica",
    "assinatura_tipo", "percentual_as",
]
ASSINATURA_TIPO = "cpu_hierarquica_v1"


def _texto(value) -> str:
    return "" if value is None else unicodedata.normalize("NFC", str(value)).strip()


def _normalizado(value) -> str:
    text = unicodedata.normalize("NFKD", _texto(value))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).casefold().split())


def _codigo(value, local: str) -> str:
    if isinstance(value, bool):
        raise ValueError(f"Código inválido em {local}")
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value != int(value):
            raise ValueError(f"Código inválido em {local}")
        value = str(int(value))
    text = _texto(value).lstrip("'")
    if not re.fullmatch(r"\d+", text) or int(text) == 0:
        raise ValueError(f"Código ausente ou inválido em {local}: {text!r}")
    return text


def _codigo_csd(value, local: str) -> str:
    text = _texto(value)
    if not text.startswith("="):
        return _codigo(value, local)
    match = re.fullmatch(r'=HYPERLINK\(.*[,;]\s*"?(\d+)"?\s*\)', text, re.I | re.S)
    if match is None:
        raise ValueError(f"Fórmula HYPERLINK não reconhecida em {local}")
    code = _codigo(match.group(1), local)
    lookups = re.findall(r'\bMATCH\(\s*"?(\d+)"?\s*[,;]', text, re.I)
    if any(_codigo(v, local) != code for v in lookups):
        raise ValueError(f"Código da fórmula diverge do MATCH em {local}")
    return code


def _decimal(value, local: str) -> Decimal:
    if isinstance(value, bool):
        raise ValueError(f"Número inválido em {local}")
    text = _texto(value).replace("\xa0", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Número inválido em {local}: {value!r}") from exc
    if not result.is_finite():
        raise ValueError(f"Número não finito em {local}")
    return result


def _cabecalho(sheet, competencia: str, tipo: str):
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = [next(iterator) for _ in range(10)]
    except StopIteration as exc:
        raise ValueError(f"Cabeçalho incompleto em {sheet.title}") from exc
    ano, mes = competencia.split("-")
    if len(header[2]) < 2 or _texto(header[2][1]) != f"{mes}/{ano}":
        raise ValueError(f"Competência divergente em {sheet.title}; esperado {competencia}")
    expected = ["grupo", "codigo da composicao"]
    expected += ["descricao", "unidade"] if tipo == "custos" else ["tipo item", "codigo do item", "descricao", "unidade", "coeficiente"]
    if [_normalizado(v) for v in header[9][:len(expected)]] != expected:
        raise ValueError(f"Cabeçalho inesperado em {sheet.title}")
    if tipo == "custos":
        if "sem desoneracao" not in _normalizado(header[1][0]):
            raise ValueError("Regime sem desoneração não confirmado em CSD")
        positions = [i for i, v in enumerate(header[8]) if _texto(v).upper() == "ES"]
        if len(positions) != 1:
            raise ValueError("Cabeçalho deve identificar ES uma única vez na linha 9")
        cost_col = positions[0]
        if (cost_col + 1 >= len(header[9])
                or _normalizado(header[9][cost_col]) != "custo (r$)"
                or _normalizado(header[9][cost_col + 1]) != "%as"):
            raise ValueError("Cabeçalho de custo/%AS do ES incompatível")
        return iterator, cost_col
    return iterator, None


def _ler_analitico(sheet, competencia: str) -> dict:
    iterator, _ = _cabecalho(sheet, competencia, "analitico")
    nodes = {}
    children = {}
    for row_number, row in enumerate(iterator, 11):
        if all(value is None for value in row):
            continue
        local = f"Analítico!B{row_number}"
        code = _codigo(row[1], local)
        tipo = _normalizado(row[2])
        description, unit = _texto(row[4]), _texto(row[5])
        if not description or not unit:
            raise ValueError(f"Descrição/unidade ausente em {local}")
        if not tipo:
            if code in nodes:
                raise ValueError(f"Código de composição duplicado em {local}: {code}")
            if row[3] is not None or row[6] is not None:
                raise ValueError(f"Linha principal inconsistente em {local}")
            nodes[code] = {"descricao": description, "unidade": unit, "linha": row_number}
        else:
            if tipo not in {"composicao", "insumo"}:
                raise ValueError(f"Tipo de item não reconhecido em {local}: {row[2]!r}")
            item = _codigo(row[3], f"Analítico!D{row_number}")
            coefficient = _decimal(row[6], f"Analítico!G{row_number}")
            if coefficient < 0:
                raise ValueError(f"Coeficiente negativo em {local}")
            children.setdefault(code, []).append({
                "tipo": tipo, "codigo": item,
                "descricao": _normalizado(description), "unidade": _normalizado(unit),
                "coeficiente": format(coefficient.normalize(), "f"),
            })
    if not nodes:
        raise ValueError("Nenhuma composição na aba Analítico")
    if any(code not in nodes for code in children):
        raise ValueError("Composição principal ausente para itens do Analítico")
    for code, node in nodes.items():
        node["itens"] = children.get(code, [])
    return nodes


def _assinaturas(nodes: dict, codes: list[str]) -> dict[str, str]:
    memo = {}
    visiting = set()

    def visit(code):
        if code in visiting:
            raise ValueError(f"Ciclo na hierarquia de composições: {code}")
        if code in memo:
            return memo[code]
        if code not in nodes:
            raise ValueError(f"Composição auxiliar ausente no Analítico: {code}")
        visiting.add(code)
        node = nodes[code]
        items = []
        for edge in node["itens"]:
            item = dict(edge)
            if edge["tipo"] == "composicao":
                item["assinatura_auxiliar"] = visit(edge["codigo"])
            items.append(item)
        payload = {
            "versao": ASSINATURA_TIPO,
            "descricao": _normalizado(node["descricao"]),
            "unidade": _normalizado(node["unidade"]),
            "itens": sorted(items, key=lambda v: json.dumps(v, sort_keys=True, ensure_ascii=False)),
        }
        serialized = json.dumps(payload, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
        memo[code] = hashlib.sha256(serialized.encode("utf-8")).hexdigest()
        visiting.remove(code)
        return memo[code]

    for code in codes:
        visit(code)
    return memo


def ler_zip_sinapi(path: str | Path, competencia: str, publicado_em: str = "") -> pd.DataFrame:
    """Retorne composições CSD do ES com custo positivo e relatório em attrs.

    `percentual_as` mantém a fração publicada (0 a 1). Hífen/nulo/zero são
    excluídos com motivo e localização em `df.attrs['relatorio']`; custos
    negativos ou textos inesperados geram erro. Não executa fórmulas.
    """
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", competencia):
        raise ValueError("Competência deve usar YYYY-MM")
    if publicado_em:
        date.fromisoformat(publicado_em)
    path = Path(path)
    raw = path.read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    expected = f"SINAPI_Referência_{competencia.replace('-', '_')}.xlsx"
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if archive.testzip() is not None:
            raise ValueError(f"CRC inválido no ZIP: {path}")
        members = [v for v in archive.namelist() if unicodedata.normalize("NFC", Path(v).name) == expected]
        if len(members) != 1:
            raise ValueError(f"Esperado um arquivo {expected}; encontrados {len(members)}")
        member = members[0]
        wb = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=False)
        try:
            if not {"CSD", "Analítico"}.issubset(wb.sheetnames):
                raise ValueError("Abas CSD e Analítico são obrigatórias")
            nodes = _ler_analitico(wb["Analítico"], competencia)
            iterator, cost_col = _cabecalho(wb["CSD"], competencia, "custos")
            rows, exclusions, seen = [], [], set()
            for row_number, row in enumerate(iterator, 11):
                if all(value is None for value in row):
                    continue
                local = f"{member}!CSD!{get_column_letter(cost_col + 1)}{row_number}"
                code = _codigo_csd(row[1], f"CSD!B{row_number}")
                if code in seen:
                    raise ValueError(f"Código duplicado em CSD: {code}, linha {row_number}")
                seen.add(code)
                if code not in nodes:
                    raise ValueError(f"Código de CSD ausente no Analítico: {code}")
                node = nodes[code]
                description, unit, group = _texto(row[2]), _texto(row[3]), _texto(row[0])
                if not description or not unit or not group:
                    raise ValueError(f"Descrição/unidade/grupo ausente em {local}")
                if (_normalizado(description) != _normalizado(node["descricao"])
                        or _normalizado(unit) != _normalizado(node["unidade"])):
                    raise ValueError(f"Correspondência CSD/Analítico divergente para {code}")
                value = row[cost_col]
                absent = _texto(value) in {"", "-", "–", "—"}
                cost = None if absent else _decimal(value, local)
                if cost is not None and cost < 0:
                    raise ValueError(f"Custo negativo em {local}")
                if absent or cost == 0:
                    exclusions.append({"codigo": code, "localizador": local,
                                       "motivo": "custo_ausente" if absent else "custo_zero"})
                    continue
                as_raw = row[cost_col + 1]
                fraction = None if _texto(as_raw) in {"", "-", "–", "—"} else _decimal(as_raw, local + " %AS")
                if fraction is not None and not 0 <= fraction <= 1:
                    raise ValueError(f"%AS fora da faixa 0–1 em {local}")
                rows.append({
                    "fonte": "sinapi_es", "competencia": competencia, "codigo": code,
                    "descricao": description, "unidade": unit, "grupo": group,
                    "custo": float(cost), "regime": "sem_desoneracao", "publicado_em": publicado_em,
                    "arquivo_sha256": sha, "localizador": local,
                    "assinatura_tipo": ASSINATURA_TIPO,
                    "percentual_as": float(fraction) if fraction is not None else float("nan"),
                })
            if not seen:
                raise ValueError("Nenhuma composição na aba CSD")
            signatures = _assinaturas(nodes, [r["codigo"] for r in rows])
            for row in rows:
                row["assinatura_tecnica"] = signatures[row["codigo"]]
        finally:
            wb.close()
    frame = pd.DataFrame(rows, columns=COLUNAS)
    frame["codigo"] = frame["codigo"].astype("string")
    frame["custo"] = frame["custo"].astype(float)
    frame["percentual_as"] = frame["percentual_as"].astype(float)
    frame.attrs["relatorio"] = {
        "fonte": "sinapi_es", "competencia": competencia, "publicado_em": publicado_em,
        "arquivo": str(path), "arquivo_sha256": sha, "zip_crc_verificado": True,
        "composicoes_analitico": len(nodes), "composicoes_lidas": len(seen),
        "composicoes_validas": len(rows), "composicoes_excluidas": len(exclusions),
        "excluidos_por_motivo": dict(Counter(v["motivo"] for v in exclusions)),
        "exclusoes": exclusions, "coluna_custo_es": get_column_letter(cost_col + 1),
        "assinatura_tipo": ASSINATURA_TIPO,
        "percentual_as_unidade": "fracao_0_a_1",
        "percentual_as_ausente": int(frame["percentual_as"].isna().sum()),
    }
    return frame


def validar_manifestos_sinapi(path: str | Path, competencia: str) -> dict:
    """Confira identidade e publicação em todos os manifestos de download.

    O catálogo de URLs/datas não prova a identidade do arquivo local.
    Esta verificação lê somente os bytes do ZIP e os JSONs, sem abrir XLSX.
    """
    path = Path(path)
    manifests = [path.parent / name for name in (
        "manifesto_coleta_regressao.json", "manifesto_download.json",
    ) if (path.parent / name).is_file()]
    if not manifests:
        raise ValueError(f"Manifesto de coleta/download obrigatório para {path.name}")
    raw = path.read_bytes()
    sha, size = hashlib.sha256(raw).hexdigest(), len(raw)
    published = set()
    for manifest_path in manifests:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if not isinstance(manifest, dict):
            raise ValueError(f"Manifesto deve ser um objeto JSON: {manifest_path}")
        if manifest.get("competencia") != competencia:
            raise ValueError(f"Competência ausente/divergente no manifesto {manifest_path}")
        if manifest.get("arquivo") != path.name:
            raise ValueError(f"Nome do arquivo diverge do manifesto {manifest_path}")
        declared_sha = manifest.get("sha256")
        if (not isinstance(declared_sha, str)
                or not re.fullmatch(r"[0-9a-fA-F]{64}", declared_sha)
                or declared_sha.lower() != sha):
            raise ValueError(f"SHA-256 ausente/divergente no manifesto {manifest_path}")
        size_keys = [key for key in ("bytes", "tamanho_bytes") if key in manifest]
        if not size_keys:
            raise ValueError(f"tamanho em bytes ausente no manifesto {manifest_path}")
        for key in size_keys:
            declared_size = manifest[key]
            if type(declared_size) is not int or declared_size != size:
                raise ValueError(f"tamanho em bytes divergente no manifesto {manifest_path}")
        for key in ("publicado_em", "publicacao_exibida_no_portal"):
            value = manifest.get(key)
            if value is None or value == "":
                continue
            if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                raise ValueError(f"Publicação inválida no manifesto {manifest_path}")
            date.fromisoformat(value)
            published.add(value)
    if len(published) > 1:
        raise ValueError(f"Publicações conflitantes nos manifestos de {competencia}")
    return {
        "arquivo": path.name, "competencia": competencia,
        "sha256": sha, "tamanho_bytes": size,
        "publicado_em": next(iter(published)) if published else "",
        "manifestos": [str(p) for p in manifests],
    }


def preparar_historico_sinapi(directory: str | Path, outputpath: str | Path) -> pd.DataFrame:
    """Prepare CSV e `.relatorio.json` de ZIPs <competência>/originais.

    Exige uma revisão por competência e manifesto de coleta/download com
    nome, competência, tamanho e SHA-256 correspondentes. Todos os manifestos
    presentes são conferidos. O catálogo pode suprir somente a data de
    publicação ausente, nunca a identidade/proveniência do ZIP.
    O relatório de exclusões é preservado também no DataFrame retornado.
    """
    directory, outputpath = Path(directory), Path(outputpath)
    paths = sorted(directory.glob("*/originais/*.zip"))
    if not paths:
        raise ValueError(f"Nenhum ZIP encontrado em {directory}/<competência>/originais")
    editions = {}
    for path in paths:
        competencia = path.parent.parent.name
        if competencia in editions:
            raise ValueError(f"Mais de uma edição/revisão em {competencia}; selecione explicitamente")
        editions[competencia] = path
    catalog_path = Path(__file__).resolve().parent.parent / "fontes_sinapi.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8")) if catalog_path.exists() else {}
    dates = {e["competencia"]: e.get("publicado_em", "") for e in catalog.get("entries", [])}
    frames, reports = [], []
    for competencia, path in editions.items():
        identity = validar_manifestos_sinapi(path, competencia)
        publication = identity["publicado_em"] or dates.get(competencia, "")
        frame = ler_zip_sinapi(path, competencia, publication)
        if frame.attrs["relatorio"]["arquivo_sha256"] != identity["sha256"]:
            raise ValueError(f"SHA-256 mudou entre validação e leitura de {path.name}")
        frame.attrs["relatorio"]["manifestos_proveniencia"] = identity["manifestos"]
        reports.append(frame.attrs["relatorio"])
        frames.append(frame)
    result = pd.concat(frames, ignore_index=True).sort_values(["competencia", "codigo"]).reset_index(drop=True)
    if result.duplicated(["fonte", "competencia", "codigo", "regime"]).any():
        raise ValueError("Código/competência duplicado no histórico")
    report = {
        "fonte": "sinapi_es", "composicoes_validas": len(result),
        "composicoes_lidas": sum(r["composicoes_lidas"] for r in reports),
        "composicoes_excluidas": sum(r["composicoes_excluidas"] for r in reports),
        "edicoes": reports,
    }
    outputpath.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", newline="", dir=outputpath.parent, delete=False) as stream:
        temporary = Path(stream.name)
        result.to_csv(stream, index=False)
    temporary.replace(outputpath)
    outputpath.with_suffix(".relatorio.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8",
    )
    result.attrs["relatorio"] = report
    return result
