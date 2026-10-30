import openpyxl

from orca_tools.excel_budget_generator import criar_planilha_orcamentaria_excel


def test_gera_pasta_de_trabalho_com_abas(tmp_path):
    saida = tmp_path / "orcamento.xlsx"
    itens = [
        {
            "item": "1.1",
            "codigo": "88316",
            "fonte": "SINAPI-ES",
            "descricao": "Servente",
            "unidade": "H",
            "quantidade": 10,
            "custo_unitario": 20.0,
            "bdi": 25.0,
        },
        {
            "item": "1.2",
            "codigo": "SRAES-CP-001",
            "fonte": "PRÓPRIA",
            "descricao": "Composição própria",
            "unidade": "m2",
            "quantidade": 5,
            "custo_unitario": 100.0,
            "bdi": 25.0,
        },
    ]
    cps = [
        {
            "codigo": "SRAES-CP-001",
            "descricao": "Composição própria",
            "unidade": "m2",
            "insumos": [
                {
                    "tipo": "MO",
                    "codigo": "88316",
                    "banco": "SINAPI",
                    "descricao": "Servente",
                    "unidade": "H",
                    "coeficiente": 0.5,
                    "preco_unitario": 20.0,
                }
            ],
        }
    ]
    caminho = criar_planilha_orcamentaria_excel(
        "Obra teste", "Vitória/ES", "2026-05", itens, cps, caminho_saida=str(saida)
    )
    wb = openpyxl.load_workbook(caminho)
    assert "01_Planilha_Orcamentaria" in wb.sheetnames
    assert len(wb.sheetnames) >= 2
