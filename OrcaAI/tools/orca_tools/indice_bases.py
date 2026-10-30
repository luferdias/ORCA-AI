"""Índice local das bases oficiais (A9 MineradorBases).

Lê os ZIPs originais de ``OrcaAI/bases/<fonte>/<competencia>/originais/`` com
os leitores já validados do módulo de ML (``orca_ml.sinapi`` e ``orca_ml.der``)
e grava um índice tabular em ``OrcaAI/knowledge/indices/<fonte>/<competencia>/``:

- ``servicos.csv``: fonte, competencia, codigo, descricao, unidade, grupo,
  custo, regime, publicado_em, arquivo_sha256, localizador.
- ``manifesto.json``: origem (ZIP e SHA-256), contagens e data de geração.

O índice é ignorado pelo Git; o manifesto pode ser versionado. A consulta por
código é O(1) depois de carregado; a busca textual é por termos normalizados.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import time
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from orca_tools.estrutura_projeto import raiz_orcaai

COLUNAS = [
    "fonte",
    "competencia",
    "codigo",
    "descricao",
    "unidade",
    "grupo",
    "custo",
    "regime",
    "publicado_em",
    "arquivo_sha256",
    "localizador",
]
FONTES = {"sinapi_es": "sinapi", "der_es": "der_es"}


def _ml_no_path() -> None:
    ml = raiz_orcaai() / "ml"
    if str(ml) not in sys.path:
        sys.path.insert(0, str(ml))


def _zip_da_competencia(fonte: str, competencia: str) -> Path:
    pasta = raiz_orcaai() / "bases" / FONTES[fonte] / competencia / "originais"
    zips = sorted(pasta.glob("*.zip"))
    if len(zips) != 1:
        raise FileNotFoundError(f"esperado um ZIP em {pasta}; encontrados {len(zips)}")
    return zips[0]


def _publicado_em(pasta: Path) -> str:
    for m in pasta.glob("manifesto*.json"):
        try:
            d = json.loads(m.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        for k in ("publicacao_exibida_no_portal", "publicado_em", "data_publicacao"):
            if d.get(k):
                return str(d[k])[:10]
    return ""


def construir(fonte: str, competencia: str, destino_raiz: Path | None = None) -> Path:
    """Constrói o índice de uma fonte/competência e devolve a pasta gerada."""
    _ml_no_path()
    origem = _zip_da_competencia(fonte, competencia)
    publicado = _publicado_em(origem.parent)
    t0 = time.time()
    if fonte == "sinapi_es":
        from orca_ml.sinapi import ler_zip_sinapi

        frame = ler_zip_sinapi(origem, competencia, publicado)
        linhas = frame[COLUNAS].to_dict("records")
        relatorio = frame.attrs.get("relatorio", {})
        insumos = _ler_insumos_sinapi(origem, competencia)
    elif fonte == "der_es":
        from orca_ml.der import parse_der_zip

        linhas = [
            {c: r.get(c, "") for c in COLUNAS}
            for r in parse_der_zip(origem, competencia, publicado)
        ]
        relatorio = {"servicos_lidos": len(linhas)}
        insumos = []
    else:
        raise ValueError(f"fonte desconhecida: {fonte}")
    destino = (
        (destino_raiz or raiz_orcaai() / "knowledge" / "indices") / fonte / competencia
    )
    destino.mkdir(parents=True, exist_ok=True)
    with open(destino / "servicos.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUNAS)
        w.writeheader()
        for r in linhas:
            w.writerow({c: r.get(c, "") for c in COLUNAS})
    if insumos:
        with open(destino / "insumos.csv", "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=COLUNAS_INSUMOS)
            w.writeheader()
            w.writerows(insumos)
    sha_csv = hashlib.sha256((destino / "servicos.csv").read_bytes()).hexdigest()
    manifesto = {
        "fonte": fonte,
        "competencia": competencia,
        "publicado_em": publicado,
        "origem_zip": str(origem.relative_to(raiz_orcaai().parent)),
        "origem_sha256": hashlib.sha256(origem.read_bytes()).hexdigest(),
        "servicos": len(linhas),
        "insumos": len(insumos),
        "servicos_csv_sha256": sha_csv,
        "gerado_em": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "duracao_s": round(time.time() - t0, 1),
        "relatorio_leitura": {k: v for k, v in relatorio.items() if k != "exclusoes"},
    }
    (destino / "manifesto.json").write_text(
        json.dumps(manifesto, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    return destino


COLUNAS_INSUMOS = [
    "fonte",
    "competencia",
    "codigo",
    "classificacao",
    "descricao",
    "unidade",
    "origem_preco",
    "preco",
    "regime",
    "localizador",
]


def _ler_insumos_sinapi(zip_path: Path, competencia: str) -> list[dict]:
    """Insumos do SINAPI (abas ISD sem desoneração e ICD com desoneração), coluna ES."""
    import io
    import unicodedata
    import zipfile

    from openpyxl import load_workbook

    esperado = f"SINAPI_Referência_{competencia.replace('-', '_')}.xlsx"
    with zipfile.ZipFile(zip_path) as z:
        nome = [
            n
            for n in z.namelist()
            if unicodedata.normalize("NFC", Path(n).name) == esperado
        ][0]
        wb = load_workbook(io.BytesIO(z.read(nome)), read_only=True, data_only=True)
    linhas: list[dict] = []
    try:
        for aba, regime in (("ISD", "sem_desoneracao"), ("ICD", "com_desoneracao")):
            if aba not in wb.sheetnames:
                continue
            ws = wb[aba]
            it = ws.iter_rows(values_only=True)
            cab = None
            for n, row in enumerate(it, 1):
                vals = [
                    str(v).replace("\n", " ").strip() if v is not None else ""
                    for v in row
                ]
                if cab is None:
                    if (
                        len(vals) > 4
                        and vals[1].lower().startswith("código do")
                        and "ES" in vals
                    ):
                        cab = vals
                        col_es = vals.index("ES")
                    continue
                codigo = vals[1]
                if not codigo or not codigo.isdigit():
                    continue
                preco = row[col_es] if col_es < len(row) else None
                linhas.append(
                    {
                        "fonte": "sinapi_es",
                        "competencia": competencia,
                        "codigo": codigo,
                        "classificacao": vals[0],
                        "descricao": vals[2],
                        "unidade": vals[3],
                        "origem_preco": vals[4],
                        "preco": "" if preco in (None, "") else preco,
                        "regime": regime,
                        "localizador": f"{aba}!B{n}",
                    }
                )
            if cab is None:
                raise ValueError(f"cabeçalho de insumos não encontrado na aba {aba}")
    finally:
        wb.close()
    return linhas


def _norm(texto: str) -> str:
    s = unicodedata.normalize("NFD", str(texto or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


@dataclass
class Indice:
    fonte: str
    competencia: str
    por_codigo: dict[str, dict]
    linhas: list[dict]

    @classmethod
    def carregar(
        cls, fonte: str, competencia: str, raiz_indices: Path | None = None
    ) -> "Indice":
        pasta = (
            (raiz_indices or raiz_orcaai() / "knowledge" / "indices")
            / fonte
            / competencia
        )
        with open(pasta / "servicos.csv", encoding="utf-8") as f:
            linhas = list(csv.DictReader(f))
        if (pasta / "insumos.csv").exists():
            with open(pasta / "insumos.csv", encoding="utf-8") as f:
                for r in csv.DictReader(f):
                    r["tipo"] = "insumo"
                    linhas.append(r)
        for r in linhas:
            r.setdefault("tipo", "servico")
            r["_busca"] = _norm(r["descricao"])
        por_codigo: dict[str, dict] = {}
        for r in linhas:
            por_codigo.setdefault(r["codigo"], r)
        return cls(fonte, competencia, por_codigo, linhas)

    def codigo(self, codigo: str) -> dict | None:
        return self.por_codigo.get(str(codigo).strip())

    def buscar(self, termos: str, limite: int = 10) -> list[dict]:
        toks = [t for t in _norm(termos).split() if len(t) > 2]
        if not toks:
            return []
        pont = []
        for r in self.linhas:
            n = sum(1 for t in toks if t in r["_busca"])
            if n:
                pont.append((n, r))
        pont.sort(key=lambda x: (-x[0], x[1]["codigo"]))
        return [r for _, r in pont[:limite]]
