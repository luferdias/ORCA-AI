import os, shutil, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT

base_dir = r'c:\Users\luis.fernando\OneDrive\04 - Meus Documentos\GitHub\Orcamento_de_Obra'
docx_path = os.path.join(base_dir, 'Despacho_SRA_Atendimento_Ata_Reuniao_SEI_63261337.docx')
pdf_path = os.path.join(base_dir, 'Despacho_SRA_Atendimento_Ata_Reuniao_SEI_63261337.pdf')

# ==========================================
# 1. GENERATE DOCX
# ==========================================
doc = docx.Document()

# Margins
for s in doc.sections:
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)

# Header Banner Table
table_header = doc.add_table(rows=2, cols=1)
table_header.alignment = WD_TABLE_ALIGNMENT.CENTER
cell_top = table_header.cell(0, 0)
cell_bot = table_header.cell(1, 0)

p_h1 = cell_top.paragraphs[0]
p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_h1 = p_h1.add_run('MINISTÉRIO DA GESTÃO E DA INOVAÇÃO EM SERVIÇOS PÚBLICOS\nSuperintendência de Administração no Espírito Santo – SRA/ES\nSetor de Engenharia / Núcleo de Infraestrutura Predial')
r_h1.font.name = 'Calibri'
r_h1.font.size = Pt(11)
r_h1.font.bold = True
r_h1.font.color.rgb = RGBColor(0, 32, 96)

p_h2 = cell_bot.paragraphs[0]
p_h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_h2 = p_h2.add_run('DESPACHO TÉCNICO – SRA/ES\nProcesso SEI nº: 10783.000206/2026-71 | Referência: Ata de Reunião SEI nº 63261337')
r_h2.font.name = 'Calibri'
r_h2.font.size = Pt(10)
r_h2.font.bold = True
r_h2.font.color.rgb = RGBColor(51, 51, 51)

doc.add_paragraph() # Spacer

# Title
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r_t = p_title.add_run('DESPACHO DE ATENDIMENTO DE ENCAMINHAMENTOS DA ATA DE REUNIÃO')
r_t.font.name = 'Calibri'
r_t.font.size = Pt(13)
r_t.font.bold = True
r_t.font.color.rgb = RGBColor(0, 32, 96)

# Meta info
p_meta = doc.add_paragraph()
r_m1 = p_meta.add_run('PARA: ')
r_m1.bold = True
p_meta.add_run('Superintendência de Administração no Espírito Santo (SRA/ES)\n')
r_m2 = p_meta.add_run('ASSUNTO: ')
r_m2.bold = True
p_meta.add_run('Atendimento integral aos requisitos orçamentários da Ata de Reunião SEI nº 63261337 – Apresentação de Demonstrativos de BDI, Justificativas Técnicas e Cronogramas Físico-Financeiros (Itens 01 a 11).\n')
r_m3 = p_meta.add_run('REFERÊNCIA: ')
r_m3.bold = True
p_meta.add_run('Processo SEI nº 10783.000206/2026-71 | Ata de Reunião CGGIP/DCD SEI nº 63261337')

# Section 1
p_s1 = doc.add_paragraph()
r_s1 = p_s1.add_run('1. CONTEXTUALIZAÇÃO E HISTÓRICO')
r_s1.font.bold = True
r_s1.font.color.rgb = RGBColor(0, 32, 96)

doc.add_paragraph(
    'Em cumprimento aos encaminhamentos registrados na Ata de Reunião de Validação do Estudo Técnico Preliminar (ETP) '
    'nº SEI 63261337, realizada entre a Coordenação-Geral de Gestão de Imóveis e Projetos (CGGIP/DCD/MGI) e a Superintendência '
    'de Administração no Espírito Santo (SRA/ES), o Setor de Engenharia submete o presente Despacho para formalizar '
    'o cumprimento integral dos requisitos técnicos e orçamentários exigidos para o prosseguimento da licitação das obras e serviços '
    'de engenharia no Edifício Sede do MGI/ES e seus Anexos.'
)

# Section 2
p_s2 = doc.add_paragraph()
r_s2 = p_s2.add_run('2. DEMONSTRATIVO DO CUMPRIMENTO DOS TRÊS REQUISITOS ORÇAMENTÁRIOS')
r_s2.font.bold = True
r_s2.font.color.rgb = RGBColor(0, 32, 96)

doc.add_paragraph(
    'A equipe de engenharia realizou a revisão e consolidação técnica dos autos, sanando rigorosamente os três apontamentos '
    'específicos de ordem orçamentária e financeira destacados no item 5 (Encaminhamentos) da supramencionada Ata de Reunião:'
)

p_req1 = doc.add_paragraph(style='List Bullet')
r_r1 = p_req1.add_run('a) Inclusão da Memória de Cálculo e Demonstrativo Analítico do BDI: ')
r_r1.bold = True
p_req1.add_run(
    'Foi estruturada a memória de cálculo completa e analítica da Bonificação e Despesas Indiretas (BDI) no percentual padrão de 25,00% '
    '(regime desonerado) para cada um dos 11 itens da contratação, detalhando a fração exata de Administração Central (3,38%), Riscos (0,85%), '
    'Seguros e Garantias (0,48%), Despesas Financeiras (0,85%), Lucro (6,00%) e tributos decorrentes do faturamento (10,45%).'
)

p_req2 = doc.add_paragraph(style='List Bullet')
r_r2 = p_req2.add_run('b) Apresentação da Fundamentação e Justificativa Técnica do BDI: ')
r_r2.bold = True
p_req2.add_run(
    'Conforme orientação da CGGIP/DCD e diretrizes do Acórdão TCU nº 2.622/2013-Plenário, da Nota Técnica SEI nº 13870/2025/MGI e da '
    'Legislação Municipal de Vitória/ES (LC 60/2002), foram elaborados pareceres técnicos circunstanciados justificando o enquadramento '
    'da taxa de BDI no 3º Quartil (25,00%), impulsionada pela alíquota teto de ISS local (5,00%) e pela complexidade executiva e '
    'normativa (NR-10, NR-18, NR-33, NR-35 e agressividade marinha de Vitória/ES).'
)

p_req3 = doc.add_paragraph(style='List Bullet')
r_r3 = p_req3.add_run('c) Anexação dos Cronogramas Físico-Financeiros: ')
r_r3.bold = True
p_req3.add_run(
    'Foram anexados os cronogramas físico-financeiros de execução para todos os 11 itens, definindo o avanço temporal das etapas, '
    'os desembolsos previstos e os critérios de medição e pagamento compatibilizados com o plano de trabalho da obra.'
)

# Section 3
p_s3 = doc.add_paragraph()
r_s3 = p_s3.add_run('3. CONSOLIDAÇÃO DOS ARTEFATOS E PUBLICAÇÃO NO SEI')
r_s3.font.bold = True
r_s3.font.color.rgb = RGBColor(0, 32, 96)

doc.add_paragraph(
    'Informa-se que a totalidade das 33 peças técnicas orçamentárias (Composição de BDI, Justificativa Técnica e Cronograma Físico-Financeiro) '
    'dos 11 itens foi unificada em um único arquivo consolidado de 33 páginas, o qual foi devidamente assinado digitalmente '
    'pelo Engenheiro Responsável Técnico e juntado aos autos do Processo SEI nº 10783.000206/2026-71 sob o título:'
)

p_file = doc.add_paragraph()
p_file.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_f = p_file.add_run('Documento SEI: BDI_Justificativa_Cronograma_Item01_ao_11_assinado.pdf')
r_f.font.bold = True
r_f.font.color.rgb = RGBColor(0, 32, 96)

# Summary table in docx
data_rows = [
    ('Item 01', 'Substituição das Chapas de Inox dos Elevadores', 'Págs. 1 a 3', 'Cumprido e Assinado'),
    ('Item 02', 'Limpeza Robotizada dos Dutos de Ar Condicionado (HVAC)', 'Págs. 4 a 6', 'Cumprido e Assinado'),
    ('Item 03', 'Lavagem da Fachada Envidraçada', 'Págs. 7 a 9', 'Cumprido e Assinado'),
    ('Item 04', 'Tratamento e Recuperação de Juntas de Dilatação', 'Págs. 10 a 12', 'Cumprido e Assinado'),
    ('Item 05', 'Pintura Anticorrosiva das Grades no Entorno', 'Págs. 13 a 15', 'Cumprido e Assinado'),
    ('Item 06', 'Substituição de Películas de Controle Solar (Insulfilm)', 'Págs. 16 a 18', 'Cumprido e Assinado'),
    ('Item 07', 'Pintura Externa do Edifício Sede MGI/ES e Anexos', 'Págs. 19 a 21', 'Cumprido e Assinado'),
    ('Item 08', 'Impermeabilização da Cisterna no Subsolo', 'Págs. 22 a 24', 'Cumprido e Assinado'),
    ('Item 09', 'Limpeza e Manutenção de Usinas Fotovoltaicas', 'Págs. 25 a 27', 'Cumprido e Assinado'),
    ('Item 10', 'Manutenção do Grupo Gerador e Subestação Elétrica', 'Págs. 28 a 30', 'Cumprido e Assinado'),
    ('Item 11', 'Reforma e Impermeabilização da Marquise do Anexo', 'Págs. 31 a 33', 'Cumprido e Assinado'),
]

table_summary = doc.add_table(rows=12, cols=4)
table_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ['Item', 'Frente de Serviço / Objeto', 'Páginas SEI', 'Status de Atendimento']
hdr_cells = table_summary.rows[0].cells
for i, h in enumerate(headers):
    hdr_cells[i].text = h
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True
    hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

for row_idx, data in enumerate(data_rows, start=1):
    row_cells = table_summary.rows[row_idx].cells
    for col_idx, text in enumerate(data):
        row_cells[col_idx].text = text
        if col_idx in [0, 2, 3]:
            row_cells[col_idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Section 4
p_s4 = doc.add_paragraph()
r_s4 = p_s4.add_run('\n4. CONCLUSÃO E ENCAMINHAMENTO')
r_s4.font.bold = True
r_s4.font.color.rgb = RGBColor(0, 32, 96)

doc.add_paragraph(
    'Diante do exposto, atesta-se o atendimento cabal aos três encaminhamentos orçamentários da Ata de Reunião SEI nº 63261337. '
    'Os autos encontram-se devidamente instruídos sob a perspectiva econômico-financeira e de planejamento. '
    'Submetem-se o presente processo e a documentação à apreciação da Superintendência (SRA/ES) para prosseguimento do feito '
    'e instrução das fases subsequentes da contratação.'
)

# Signature block
p_sig = doc.add_paragraph()
p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sig.add_run('\n____________________________________________________\n').bold = False
p_sig.add_run('Luis Fernando Dias\n').bold = True
p_sig.add_run('Engenheiro Civil – CREA/ES\nSetor de Engenharia — SRA/ES / MGI').italic = True

doc.save(docx_path)
print('DOCX Despacho generated successfully!')


# ==========================================
# 2. GENERATE REPORTLAB PDF
# ==========================================
doc_pdf = SimpleDocTemplate(
    pdf_path,
    pagesize=A4,
    rightMargin=36,
    leftMargin=36,
    topMargin=36,
    bottomMargin=36
)

styles = getSampleStyleSheet()

style_banner_title = ParagraphStyle(
    'BannerTitle',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=10,
    leading=12,
    textColor=colors.HexColor('#002060'),
    alignment=TA_CENTER
)

style_banner_sub = ParagraphStyle(
    'BannerSub',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=8.5,
    leading=11,
    textColor=colors.HexColor('#333333'),
    alignment=TA_CENTER
)

style_title = ParagraphStyle(
    'DocTitle',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=10.5,
    leading=13,
    textColor=colors.HexColor('#002060'),
    spaceBefore=6,
    spaceAfter=4
)

style_heading = ParagraphStyle(
    'HeadingSection',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=9,
    leading=11,
    textColor=colors.HexColor('#002060'),
    spaceBefore=6,
    spaceAfter=3
)

style_body = ParagraphStyle(
    'BodyJustified',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=8,
    leading=10.5,
    textColor=colors.black,
    alignment=TA_JUSTIFY,
    spaceBefore=2,
    spaceAfter=3
)

style_table_header = ParagraphStyle(
    'TableHeader',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=7,
    leading=9,
    textColor=colors.HexColor('#002060'),
    alignment=TA_CENTER
)

style_table_cell = ParagraphStyle(
    'TableCell',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=6.5,
    leading=8.5,
    textColor=colors.black,
    alignment=TA_LEFT
)

style_table_cell_center = ParagraphStyle(
    'TableCellCenter',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=6.5,
    leading=8.5,
    textColor=colors.black,
    alignment=TA_CENTER
)

story_pdf = []

# Banner
b_p1 = Paragraph('MINISTÉRIO DA GESTÃO E DA INOVAÇÃO EM SERVIÇOS PÚBLICOS<br/>Superintendência de Administração no Espírito Santo – SRA/ES<br/>Setor de Engenharia / Núcleo de Infraestrutura Predial', style_banner_title)
b_p2 = Paragraph('DESPACHO TÉCNICO – SRA/ES<br/>Processo SEI nº: 10783.000206/2026-71 | Referência: Ata de Reunião SEI nº 63261337', style_banner_sub)

banner_table = Table([[b_p1], [b_p2]], colWidths=[520])
banner_table.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F2F4F7')),
    ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#002060')),
    ('PADDING', (0,0), (-1,-1), 4),
    ('ALIGN', (0,0), (-1,-1), 'CENTER'),
]))

story_pdf.append(banner_table)
story_pdf.append(Spacer(1, 4))

# Title
story_pdf.append(Paragraph('DESPACHO DE ATENDIMENTO DE ENCAMINHAMENTOS DA ATA DE REUNIÃO SEI Nº 63261337', style_title))

# Meta
meta_text = (
    '<b>PARA:</b> Superintendência de Administração no Espírito Santo (SRA/ES)<br/>'
    '<b>ASSUNTO:</b> Atendimento integral aos requisitos orçamentários da Ata de Reunião SEI nº 63261337 – Apresentação de Demonstrativos de BDI, Justificativas Técnicas e Cronogramas Físico-Financeiros (Itens 01 a 11).<br/>'
    '<b>REFERÊNCIA:</b> Processo SEI nº 10783.000206/2026-71 | Ata de Reunião CGGIP/DCD SEI nº 63261337'
)
story_pdf.append(Paragraph(meta_text, style_body))
story_pdf.append(Spacer(1, 3))

# Section 1
story_pdf.append(Paragraph('1. CONTEXTUALIZAÇÃO E HISTÓRICO', style_heading))
sec1_pdf = (
    'Em cumprimento aos encaminhamentos registrados na Ata de Reunião de Validação do Estudo Técnico Preliminar (ETP) '
    'nº SEI 63261337, realizada entre a Coordenação-Geral de Gestão de Imóveis e Projetos (CGGIP/DCD/MGI) e a Superintendência '
    'de Administração no Espírito Santo (SRA/ES), o Setor de Engenharia submete o presente Despacho para formalizar '
    'o cumprimento integral dos requisitos técnicos e orçamentários exigidos para o prosseguimento da licitação das obras e serviços '
    'de engenharia no Edifício Sede do MGI/ES e seus Anexos.'
)
story_pdf.append(Paragraph(sec1_pdf, style_body))

# Section 2
story_pdf.append(Paragraph('2. DEMONSTRATIVO DO CUMPRIMENTO DOS TRÊS REQUISITOS ORÇAMENTÁRIOS', style_heading))
sec2_pdf_intro = (
    'A equipe de engenharia realizou a revisão e consolidação técnica dos autos, sanando rigorosamente os três apontamentos '
    'específicos de ordem orçamentária e financeira destacados no item 5 (Encaminhamentos) da supramencionada Ata de Reunião:'
)
story_pdf.append(Paragraph(sec2_pdf_intro, style_body))

req1_pdf = (
    '• <b>a) Inclusão da Memória de Cálculo e Demonstrativo Analítico do BDI:</b> Foi estruturada a memória de cálculo completa e analítica '
    'da Bonificação e Despesas Indiretas (BDI) no percentual padrão de 25,00% (regime desonerado) para cada um dos 11 itens da contratação, '
    'detalhando a fração exata de Administração Central (3,38%), Riscos (0,85%), Seguros e Garantias (0,48%), Despesas Financeiras (0,85%), '
    'Lucro (6,00%) e tributos decorrentes do faturamento (10,45%).'
)
req2_pdf = (
    '• <b>b) Apresentação da Fundamentação e Justificativa Técnica do BDI:</b> Conforme orientação da CGGIP/DCD e diretrizes do Acórdão TCU '
    'nº 2.622/2013-Plenário, da Nota Técnica SEI nº 13870/2025/MGI e da Legislação Municipal de Vitória/ES (LC 60/2002), foram elaborados pareceres '
    'técnicos circunstanciados justificando o enquadramento da taxa de BDI no 3º Quartil (25,00%), impulsionada pela alíquota teto de ISS local (5,00%) '
    'e pela complexidade executiva e normativa (NR-10, NR-18, NR-33, NR-35 e agressividade marinha de Vitória/ES).'
)
req3_pdf = (
    '• <b>c) Anexação dos Cronogramas Físico-Financeiros:</b> Foram anexados os cronogramas físico-financeiros de execução para todos os 11 itens, '
    'definindo o avanço temporal das etapas, os desembolsos previstos e os critérios de medição e pagamento compatibilizados com o plano de trabalho.'
)
story_pdf.append(Paragraph(req1_pdf, style_body))
story_pdf.append(Paragraph(req2_pdf, style_body))
story_pdf.append(Paragraph(req3_pdf, style_body))

# Section 3
story_pdf.append(Paragraph('3. CONSOLIDAÇÃO DOS ARTEFATOS E PUBLICAÇÃO NO SEI', style_heading))
sec3_pdf = (
    'Informa-se que a totalidade das 33 peças técnicas orçamentárias (Composição de BDI, Justificativa Técnica e Cronograma Físico-Financeiro) '
    'dos 11 itens foi unificada em um único arquivo consolidado de 33 páginas, o qual foi devidamente assinado digitalmente '
    'pelo Engenheiro Responsável Técnico e juntado aos autos do Processo SEI nº 10783.000206/2026-71 sob o título: '
    '<b><font color="#002060">BDI_Justificativa_Cronograma_Item01_ao_11_assinado.pdf</font></b>.'
)
story_pdf.append(Paragraph(sec3_pdf, style_body))

# Table in pdf
table_pdf_data = [
    [
        Paragraph('Item', style_table_header),
        Paragraph('Frente de Serviço / Objeto', style_table_header),
        Paragraph('Páginas SEI', style_table_header),
        Paragraph('Status de Atendimento', style_table_header)
    ]
]

for row in data_rows:
    table_pdf_data.append([
        Paragraph(row[0], style_table_cell_center),
        Paragraph(row[1], style_table_cell),
        Paragraph(row[2], style_table_cell_center),
        Paragraph(row[3], style_table_cell_center)
    ])

summary_pdf_table = Table(table_pdf_data, colWidths=[45, 255, 95, 125])
summary_pdf_table.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EBF1F5')),
    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#B0C4DE')),
    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
    ('PADDING', (0,0), (-1,-1), 2),
]))

story_pdf.append(summary_pdf_table)
story_pdf.append(Spacer(1, 4))

# Section 4
story_pdf.append(Paragraph('4. CONCLUSÃO E ENCAMINHAMENTO', style_heading))
sec4_pdf = (
    'Diante do exposto, atesta-se o atendimento cabal aos três encaminhamentos orçamentários da Ata de Reunião SEI nº 63261337. '
    'Os autos encontram-se devidamente instruídos sob a perspectiva econômico-financeira e de planejamento. '
    'Submetem-se o presente processo e a documentação à apreciação da Superintendência (SRA/ES) para prosseguimento do feito '
    'e instrução das fases subsequentes da contratação.'
)
story_pdf.append(Paragraph(sec4_pdf, style_body))

story_pdf.append(Spacer(1, 8))

style_sig_line = ParagraphStyle('SigLine', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, alignment=TA_CENTER)
style_sig_title = ParagraphStyle('SigTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, alignment=TA_CENTER)
style_sig_sub = ParagraphStyle('SigSub', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=7.5, alignment=TA_CENTER)

story_pdf.append(Paragraph('____________________________________________________', style_sig_line))
story_pdf.append(Paragraph('Luis Fernando Dias', style_sig_title))
story_pdf.append(Paragraph('Engenheiro Civil – CREA/ES<br/>Setor de Engenharia — SRA/ES / MGI', style_sig_sub))

doc_pdf.build(story_pdf)
print('PDF Despacho generated successfully!')
