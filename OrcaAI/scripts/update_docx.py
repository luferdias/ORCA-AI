from docx import Document

def main():
    doc_path = 'Termo_Justificativas_Consolidado.docx'
    doc = Document(doc_path)
    
    doc.add_page_break()
    doc.add_heading('--- JUSTIFICATIVAS PREENCHIDAS (GERADAS) ---', level=1)
    
    # Item 10
    doc.add_heading('Item 10: Justificativas técnicas de BDI', level=2)
    doc.add_paragraph('- Administração Central: Justifica-se pela multiplicidade de serviços simultâneos em três imóveis distintos (Sede, Vila Rubim e Princesa Isabel), exigindo maior coordenação gerencial e logística.')
    doc.add_paragraph('- Seguro e Garantia: Decorre da exigência de apólices robustas para cobrir riscos de serviços complexos (altura, impermeabilização e subestações/geradores).')
    doc.add_paragraph('- Risco: Compatível com as incertezas, patologias ocultas e restrições operacionais de manutenções em edificações ocupadas.')
    doc.add_paragraph('- Despesa Financeira: Reflete o custo de capital no intervalo entre o desembolso da empresa e o pagamento pela Administração.')
    doc.add_paragraph('- Lucro: Adequado à complexidade e riscos do escopo multifacetado, garantindo a viabilidade do contrato conforme o Acórdão 2.622/2013-TCU.')
    
    # Item 14
    doc.add_heading('Item 14: Capacidade Técnico-Operacional e Técnico-Profissional', level=2)
    p = doc.add_paragraph()
    p.add_run('Capacidade Técnico-Operacional (da Empresa):\n').bold = True
    p.add_run('Atestados comprovando execução prévia de:')
    doc.add_paragraph('1. Manutenção predial preventiva/corretiva;')
    doc.add_paragraph('2. Manutenção em grupos geradores;')
    doc.add_paragraph('3. Serviços de limpeza/manutenção em altura (fachadas);')
    doc.add_paragraph('4. Impermeabilização estrutural.')

    p2 = doc.add_paragraph()
    p2.add_run('Capacidade Técnico-Profissional (do Responsável Técnico):\n').bold = True
    p2.add_run('Certidão de Acervo Técnico (CAT) atestando responsabilidade por:')
    doc.add_paragraph('1. Direção ou execução de reforma/manutenção predial;')
    doc.add_paragraph('2. Manutenção de instalações elétricas e/ou grupos geradores.')

    # Item 20
    doc.add_heading('Item 20: Necessidade de garantia da execução', level=2)
    p3 = doc.add_paragraph()
    p3.add_run('Sim. ').bold = True
    p3.add_run('Exige-se garantia de 5% (art. 96 da Lei 14.133/2021) para assegurar o fiel cumprimento do contrato, cobrir eventuais multas e mitigar prejuízos à Administração, especialmente frente aos riscos de serviços em altura, elétricos e de impermeabilização.')

    # Item 21
    doc.add_heading('Item 21: Requisitos de sustentabilidade', level=2)
    doc.add_paragraph('1. Cumprimento do Plano de Gerenciamento de Resíduos da Construção Civil (PGRCC) para destinação correta de entulhos (Resolução CONAMA 307/2002);')
    doc.add_paragraph('2. Priorização de materiais de baixo impacto (tintas sem solvente/baixo VOC e limpadores biodegradáveis);')
    doc.add_paragraph('3. Uso de equipamentos de baixo consumo de água e energia;')
    doc.add_paragraph('4. Logística reversa no descarte de óleos lubrificantes, filtros, pilhas e baterias do gerador.')
    
    doc.save(doc_path)
    print("Documento atualizado com sucesso!")

if __name__ == '__main__':
    main()
