#!/usr/bin/env python3
"""Validate SINAPI and IOPES codes from local spreadsheets and CSV files."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None


FIELD_ALIASES = {
    "system": (
        "sistema",
        "origem",
        "fonte",
        "base",
        "tabela",
        "referencia",
        "banco",
    ),
    "code": (
        "codigo",
        "cod",
        "item",
        "codigo item",
        "codigo composicao",
        "codigo insumo",
    ),
    "description": (
        "descricao",
        "descrição",
        "especificacao",
        "especificação",
        "servico",
        "serviço",
        "insumo",
        "item descricao",
    ),
    "unit": ("unidade", "un", "und"),
    "price": ("preco", "preço", "valor", "custo", "preco unitario", "valor unitario"),
}

SYSTEM_NAMES = ("sinapi", "iopes")


@dataclass
class RowRecord:
    source_path: str
    sheet: Optional[str]
    row_number: int
    system: Optional[str]
    code_raw: str
    code_norm: str
    description: str
    unit: str
    price: str
    row: Dict[str, str]


def normalize_header(value: str) -> str:
    collapsed = re.sub(r"\s+", " ", value.strip().lower())
    return (
        collapsed.replace("ç", "c")
        .replace("ã", "a")
        .replace("á", "a")
        .replace("à", "a")
        .replace("â", "a")
        .replace("é", "e")
        .replace("ê", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ô", "o")
        .replace("õ", "o")
        .replace("ú", "u")
    )


def normalize_code(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9]", "", (value or "").upper())


def normalize_system(value: Optional[str], fallback_name: str = "") -> Optional[str]:
    candidates = [value or "", fallback_name]
    for candidate in candidates:
        text = normalize_header(candidate)
        for system in SYSTEM_NAMES:
            if system in text:
                return system
    return None


def first_non_empty(values: Iterable[str]) -> str:
    for value in values:
        if value and str(value).strip():
            return str(value).strip()
    return ""


def detect_columns(headers: Sequence[str]) -> Dict[str, Optional[str]]:
    normalized = {header: normalize_header(header) for header in headers}
    mapping: Dict[str, Optional[str]] = {}
    for field, aliases in FIELD_ALIASES.items():
        chosen = None
        for header, normalized_header in normalized.items():
            if normalized_header in aliases:
                chosen = header
                break
        if chosen is None:
            for header, normalized_header in normalized.items():
                if any(alias in normalized_header for alias in aliases):
                    chosen = header
                    break
        mapping[field] = chosen
    return mapping


def parse_path_and_sheet(raw: str) -> Tuple[Path, Optional[str]]:
    if "#" in raw:
        path_str, sheet = raw.split("#", 1)
        return Path(path_str), sheet
    return Path(raw), None


def load_rows(path_arg: str) -> Tuple[List[Dict[str, str]], Optional[str], str]:
    path, requested_sheet = parse_path_and_sheet(path_arg)
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            return list(reader), None, str(path)
    if suffix == ".json":
        with path.open("r", encoding="utf-8") as handle:
            data = json.load(handle)
        if not isinstance(data, list):
            raise ValueError(f"{path} deve conter uma lista JSON de objetos.")
        rows = [dict((str(k), "" if v is None else str(v)) for k, v in item.items()) for item in data]
        return rows, None, str(path)
    if suffix in {".xlsx", ".xlsm"}:
        if load_workbook is None:
            raise RuntimeError(
                "openpyxl nao esta disponivel. Instale a dependencia ou exporte o arquivo para CSV."
            )
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet_name = requested_sheet or workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"A aba '{sheet_name}' nao existe em {path}.")
        sheet = workbook[sheet_name]
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return [], sheet_name, str(path)
        headers = ["" if value is None else str(value) for value in values[0]]
        rows = []
        for raw_row in values[1:]:
            row = {}
            for index, header in enumerate(headers):
                if not header:
                    continue
                cell_value = raw_row[index] if index < len(raw_row) else ""
                row[header] = "" if cell_value is None else str(cell_value)
            rows.append(row)
        return rows, sheet_name, str(path)
    raise ValueError(f"Formato nao suportado: {path.suffix or path.name}")


def records_from_rows(path_arg: str, fallback_system: Optional[str] = None) -> List[RowRecord]:
    rows, sheet_name, source_path = load_rows(path_arg)
    if not rows:
        return []
    columns = detect_columns(list(rows[0].keys()))
    code_column = columns["code"]
    if code_column is None:
        raise ValueError(f"Nenhuma coluna de codigo foi detectada em {source_path}.")

    fallback_from_name = normalize_system(fallback_system, Path(source_path).name)
    records: List[RowRecord] = []
    for index, row in enumerate(rows, start=2):
        code_raw = first_non_empty([row.get(code_column, "")])
        code_norm = normalize_code(code_raw)
        if not code_norm:
            continue
        system_value = row.get(columns["system"], "") if columns["system"] else ""
        description = row.get(columns["description"], "") if columns["description"] else ""
        unit = row.get(columns["unit"], "") if columns["unit"] else ""
        price = row.get(columns["price"], "") if columns["price"] else ""
        records.append(
            RowRecord(
                source_path=source_path,
                sheet=sheet_name,
                row_number=index,
                system=normalize_system(system_value, fallback_from_name or ""),
                code_raw=code_raw,
                code_norm=code_norm,
                description=description.strip(),
                unit=unit.strip(),
                price=price.strip(),
                row={str(k): "" if v is None else str(v) for k, v in row.items()},
            )
        )
    return records


def similarity_token_overlap(left: str, right: str) -> float:
    left_tokens = {token for token in re.split(r"\W+", normalize_header(left)) if token}
    right_tokens = {token for token in re.split(r"\W+", normalize_header(right)) if token}
    if not left_tokens or not right_tokens:
        return 0.0
    intersection = len(left_tokens & right_tokens)
    union = len(left_tokens | right_tokens)
    return intersection / union if union else 0.0


def compare_records(source_records: Sequence[RowRecord], catalogs: Dict[str, Dict[str, RowRecord]]) -> Dict[str, object]:
    results = []
    summary = {
        "checked": 0,
        "matched": 0,
        "missing": 0,
        "description_mismatch": 0,
        "unit_mismatch": 0,
        "unknown_system": 0,
    }

    for record in source_records:
        summary["checked"] += 1
        result = {
            "source_path": record.source_path,
            "sheet": record.sheet,
            "row_number": record.row_number,
            "system": record.system or "unknown",
            "code": record.code_raw,
            "normalized_code": record.code_norm,
            "status": "ok",
            "issues": [],
        }

        if not record.system or record.system not in catalogs:
            summary["unknown_system"] += 1
            result["status"] = "unknown-system"
            result["issues"].append("Sistema nao identificado para este codigo.")
            results.append(result)
            continue

        catalog_row = catalogs[record.system].get(record.code_norm)
        if catalog_row is None:
            summary["missing"] += 1
            result["status"] = "missing"
            result["issues"].append(f"Codigo nao encontrado na base {record.system.upper()}.")
            results.append(result)
            continue

        mismatch = False
        overlap = similarity_token_overlap(record.description, catalog_row.description)
        if record.description and catalog_row.description and overlap < 0.45:
            mismatch = True
            summary["description_mismatch"] += 1
            result["issues"].append(
                f"Descricao divergente. Orcamento: '{record.description}' | Base: '{catalog_row.description}'."
            )

        if record.unit and catalog_row.unit and normalize_header(record.unit) != normalize_header(catalog_row.unit):
            mismatch = True
            summary["unit_mismatch"] += 1
            result["issues"].append(
                f"Unidade divergente. Orcamento: '{record.unit}' | Base: '{catalog_row.unit}'."
            )

        if mismatch:
            result["status"] = "mismatch"
        else:
            summary["matched"] += 1
            result["issues"].append("Codigo localizado e consistente com a base local.")

        results.append(result)

    return {"summary": summary, "results": results}


def build_catalog_map(catalog_args: Sequence[str]) -> Dict[str, Dict[str, RowRecord]]:
    catalog_map: Dict[str, Dict[str, RowRecord]] = {}
    for catalog_arg in catalog_args:
        if "=" not in catalog_arg:
            raise ValueError("Use --catalog sinapi=arquivo ou --catalog iopes=arquivo.")
        system_name, raw_path = catalog_arg.split("=", 1)
        system = normalize_system(system_name)
        if system not in SYSTEM_NAMES:
            raise ValueError(f"Sistema invalido em --catalog: {system_name}")
        catalog_rows = records_from_rows(raw_path, fallback_system=system)
        catalog_map.setdefault(system, {})
        for row in catalog_rows:
            catalog_map[system][row.code_norm] = row
    return catalog_map


def format_report(report: Dict[str, object]) -> str:
    summary = report["summary"]
    results = report["results"]
    lines = [
        "Resumo da validacao",
        f"- Itens verificados: {summary['checked']}",
        f"- Itens consistentes: {summary['matched']}",
        f"- Codigos ausentes: {summary['missing']}",
        f"- Divergencias de descricao: {summary['description_mismatch']}",
        f"- Divergencias de unidade: {summary['unit_mismatch']}",
        f"- Sistema nao identificado: {summary['unknown_system']}",
        "",
        "Detalhes",
    ]
    if not results:
        lines.append("- Nenhum codigo encontrado para validar.")
        return "\n".join(lines)

    for item in results:
        location = f"{item['source_path']}:{item['row_number']}"
        if item["sheet"]:
            location += f" ({item['sheet']})"
        issues = " | ".join(item["issues"])
        lines.append(
            f"- [{item['status']}] {item['system'].upper()} {item['code']} em {location}: {issues}"
        )
    return "\n".join(lines)


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Valida codigos SINAPI e IOPES em orcamentos locais."
    )
    parser.add_argument(
        "--source",
        action="append",
        required=True,
        help="Arquivo de orcamento/lista de codigos (.csv, .tsv, .json, .xlsx). Use caminho#Aba para selecionar uma aba especifica.",
    )
    parser.add_argument(
        "--catalog",
        action="append",
        required=True,
        help="Arquivo base no formato sistema=caminho. Ex.: sinapi=sinapi.xlsx ou iopes=iopes.csv",
    )
    parser.add_argument(
        "--json-out",
        help="Caminho opcional para salvar o relatorio em JSON.",
    )
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    catalog_map = build_catalog_map(args.catalog)
    source_records: List[RowRecord] = []
    for source_arg in args.source:
        source_records.extend(records_from_rows(source_arg))
    report = compare_records(source_records, catalog_map)
    print(format_report(report))
    if args.json_out:
        output_path = Path(args.json_out)
        output_path.write_text(json.dumps(report, indent=2, ensure_ascii=True), encoding="utf-8")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # pragma: no cover - CLI safety
        print(f"Erro: {exc}", file=sys.stderr)
        raise SystemExit(1)
